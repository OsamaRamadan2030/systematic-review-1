import sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.section import WD_ORIENT
from docxlib import para, table, clear_body_after, add_runs

SRC, FIG, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
doc = Document(SRC)

# Keep the title and the Introduction exactly as supplied; rebuild everything from "2 Methods" onward.
idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "2 Methods")
body_children = [c for c in doc.element.body.iterchildren()]
methods_el = doc.paragraphs[idx]._p
clear_body_after(doc, keep_first=body_children.index(methods_el))

H1 = lambda t: para(doc, t, "Heading 1")
H2 = lambda t: para(doc, t, "Heading 2")
H3 = lambda t: para(doc, t, "Heading 3")
P = lambda t: para(doc, t)

H1("2 Methods")

# ---------------------------------------------------------------- 2.1
H2("2.1 Design, reporting and protocol history")
P("This was a concurrent, controlled, three-arm quasi-experimental study with baseline and post-instruction assessments. "
  "Intact laboratory sections were allocated nonrandomly to clinically reviewed AI-generated mind maps, student-generated "
  "mind maps or structured case-analysis worksheets without maps. Sections were the allocation units, groups within "
  "sections were the teaching units and students were the outcome-measurement and analysis units. All arms received the "
  "same flipped preparation, curriculum, contact time, teaching cases, simulation practice and feedback criteria; the arms "
  "differed only in the source of the case representation and the learner activity performed with it (Table 1). Because "
  "map source and learner activity varied together, the comparison evaluated instructional strategies rather than an "
  "isolated effect of AI generation. Reporting follows the TREND statement for nonrandomised evaluations, TIDieR and GREET "
  "for intervention description and the reporting guidelines for health care simulation research (Des Jarlais et al., 2004; "
  "Hoffmann et al., 2014; Cheng et al., 2016; Phillips et al., 2016). Table 4 gives the sequence of study procedures.")
P("The study was not registered in a public registry. The protocol and statistical analysis plan (version 1) were dated "
  "15 November 2025, before ethics approval and data collection. Version 2 of the analysis plan was dated 30 June 2026 "
  "[[and was finalised after rating was complete and before database lock and release of the strategy code key on "
  "30 June 2026; state the recorded order and times]]. [[List each difference between versions 1 and 2 with its reason.]] "
  "Analyses are described as prespecified only when they appear in version 1; analyses introduced in version 2 are "
  "identified as such. This report addresses clinical reasoning, knowledge, procedural performance and implementation. "
  "Additional exploratory process measures [[name them]] were collected in the broader study and are not reported here; "
  "[[none was a protocol primary or secondary outcome / if one was, state its omission and the reason]], and the decision "
  "to report them separately was made [[before unmasking; give date]]. The protocol, both analysis-plan versions and the "
  "amendment log are available from the corresponding author on reasonable request.")

# ---------------------------------------------------------------- 2.2
H2("2.2 Setting and participants")
P("The study was embedded in the compulsory third-year Pediatric Nursing course in the Pediatric Clinical Skills "
  "Laboratory, Faculty of Nursing, Tanta University, Egypt, during the second semester of the 2025–2026 academic year. "
  "Before the study, the faculty had timetabled the 450 enrolled students into 15 laboratory sections of 30 students. "
  "Each section contained six groups of five students, [[formed by the faculty before allocation by roster order and "
  "kept unchanged throughout the unit; confirm the rule]]. Five facilitators [[academic rank and years of pediatric "
  "teaching experience]] each taught three sections. English was the language of the instructional materials and "
  "knowledge test; students could explain their reasoning in Arabic or English during assessment.")
P("Students enrolled in the course for the first time were eligible. Students repeating the course or reporting "
  "previous structured neurotrauma simulation training outside the curriculum were to be excluded; [[no enrolled student "
  "met an exclusion criterion]], and all 450 were eligible. Recruitment took place on [[date]], after ethics approval and "
  "before baseline assessment. It was conducted by a researcher with no teaching or grading role, before allocation was "
  "determined. Written information was provided in [[Arabic and English]]. Students were informed that the teaching was "
  "part of the curriculum irrespective of participation and that declining or withdrawing would not affect grades, standing "
  "or access to teaching. All 450 students provided written consent. Attendance did not determine research eligibility.")
P("Baseline variables were age, sex, previous-year grade point average (GPA), English course grade, number of previous "
  "simulation sessions, previous mind-mapping use, previous generative AI use and previous exposure to children with head "
  "injury during clinical placements. GPA and English course grade were obtained from authorised faculty records; the "
  "remaining variables were collected by baseline questionnaire. Preparation engagement and attendance were recorded "
  "separately as implementation variables.")

# ---------------------------------------------------------------- 2.3
H2("2.3 Sample size and precision")
P("The eligible cohort was recruited by census, so enrolment and the fixed number of sections determined the sample size. "
  "Precision was planned for the principal contrast between the two mapping strategies on the primary outcome. The planning "
  "assumptions were a 2.5-point difference on the 0–24 reasoning score [[state why 2.5 points was judged the smallest "
  "educationally important difference]], a standard deviation of 5 (standardised difference 0.50), 10% loss to follow-up, a "
  "section intraclass correlation of 0.03, an additional group-within-section intraclass correlation of 0.05 and a "
  "baseline–follow-up correlation of 0.50. For sections of 27 analysed students in groups averaging 4.5, these assumptions "
  "imply a design effect of approximately 1.96, calculated as 1 + (4.5 − 1)(0.03 + 0.05) + (27 − 4.5)(0.03), and a 25% "
  "reduction in residual variance from baseline adjustment (1 − 0.50²). With 135 analysed students per arm, they imply "
  "approximately 84% power for the principal contrast at a two-sided α of 0.05. This calculation used a t reference "
  "distribution with eight section-level degrees of freedom (15 sections minus three strategy and four facilitator "
  "parameters). The corresponding minimum detectable difference at 80% power was approximately 0.47 standard deviations "
  "(2.4 points). [[If the protocol contains the original planning calculation, report it here instead.]] Precision "
  "achieved is conveyed by confidence intervals; post hoc power was not calculated.")

# ---------------------------------------------------------------- 2.4
H2("2.4 Nonrandom allocation")
P("An independent methodologist with no teaching or assessment role applied a deterministic allocation rule on "
  "12 February 2026, after baseline assessment and before orientation. Allocation was constrained within facilitator: "
  "each facilitator's three sections received one strategy each. This produced five sections and 150 students per arm and "
  "prevented confounding of facilitator with strategy. The constraint permits 6⁵ = 7,776 allocations. Admissible allocations "
  "balanced teaching position in the timetable and limited the between-arm difference in afternoon sections to one "
  "[[number of admissible allocations]]. From the admissible set, the selected allocation minimised the sum of absolute "
  "pairwise standardised differences in section means for GPA, baseline knowledge, English course grade, previous "
  "simulation sessions and previous mapping use; remaining ties were resolved by timetable order. The rule contained no "
  "random element. Baseline reasoning recordings were rated only after follow-up and could not inform allocation. Students "
  "could not select or change their strategy.")
P("Random procedures were used for assessment-case assignment and order, knowledge-test form sequence, rating order and "
  "the selection of double-rating and monitoring samples, using [[software and function; seeds retained in the study "
  "records]]. Assessment scheduling included all three strategies under each facilitator on each assessment day, and each "
  "student's teaching-to-assessment interval was recorded.")

# ---------------------------------------------------------------- 2.5
H2("2.5 Educational framework and clinical content")
P("Generative learning informed the comparison between students organising case information themselves and interpreting "
  "and revising a supplied representation (Fiorella and Mayer, 2016). Cognitive load theory provided a complementary "
  "rationale for initial structure (Sweller et al., 2019). These mechanisms were not tested separately. The Clinical "
  "Reasoning Cycle supplied eight common headings for all activities and for assessment: situation, cues, interpretation, "
  "problems, goals, actions, evaluation and reflection (Levett-Jones et al., 2010).")
P("The unit covered the head and brain injury component of pediatric neurotrauma in children aged six months to 12 years. "
  "Content addressed developmentally appropriate neurological assessment, recognition of deterioration, oxygenation and "
  "circulation, prioritisation, escalation, reassessment and family communication. Materials were reviewed against "
  "pediatric head-injury and severe traumatic brain injury guidance, and actions were framed within the nursing "
  "responsibilities represented in the cases (Kochanek et al., 2019; National Institute for Health and Care Excellence, "
  "2023; Chan et al., 2025). [[Number]] learning objectives were mapped to the eight reasoning headings and to the "
  "assessment blueprints of all three outcome instruments. Scenarios were designed using the Healthcare Simulation "
  "Standards of Best Practice for simulation design (Watts et al., 2021) [[confirm]].")

# ---------------------------------------------------------------- 2.6
H2("2.6 Cases and AI-generated maps")
P("Twenty-one fictional cases were developed: one dehydration orientation case for practising the arm-specific procedure, "
  "four neurotrauma index cases (one per teaching session), eight teaching-simulation variants (two per session) and eight "
  "assessment cases (four young-child and four school-age). Assessment cases were withheld from students, teaching "
  "facilitators and simulation instructors, and access was restricted to the assessment team. Each session's shared "
  "preparation package comprised narrated slides, a reading guide and an ungraded self-check quiz requiring approximately "
  "60 minutes. It was released [[through the learning-management system, number of days before each session]] and "
  "contained no strategy-specific map or worksheet.")
P("Two educators [[roles; state that neither taught a study section]] generated three candidate outlines for each index "
  "and orientation case (15 outputs) using GPT-4o through a web interface on 5–8 January 2026 with default generation "
  "settings. Each generation used a separate temporary conversation with memory and custom instructions disabled. Inputs "
  "comprised the fictional case, learning objectives, the eight reasoning headings, a team-written guideline summary and "
  "formatting constraints: up to three hierarchical levels and short keyword nodes. The complete prompt text was retained. "
  "No identifiable patient or student data were entered. The educators independently rated each candidate for clinical "
  "quality [[criteria and scale]], selected the highest-scoring outline for each case and resolved ties by discussion. "
  "Selected outlines were formatted as hierarchical mind maps; formatting changes were recorded separately from content "
  "changes.")
P("Three independent clinicians, a pediatric nursing professor, a pediatric critical care nurse specialist and a pediatric "
  "neurosurgeon, reviewed the selected maps for factual accuracy, age appropriateness, nursing scope, prioritisation, "
  "omissions and clarity. Each had at least ten years' relevant experience and no teaching or outcome-rating role. They "
  "independently recommended retention, modification, deletion or addition for each node and then resolved disagreements "
  "by consensus. Review occurred on 11–15 January, and final map content was frozen on 19 January 2026. The change log "
  "classified every final node as unchanged AI content, modified AI content or expert addition, and these proportions are "
  "reported for each map to characterise the AI-derived content actually delivered. [[The same panel reviewed the "
  "facilitator feedback guides used in all three arms; confirm.]] Maps were labelled as prepared with AI assistance and "
  "clinically reviewed. Students received the reviewed maps and did not interact with AI during the intervention.")

# ---------------------------------------------------------------- 2.7
H2("2.7 Instructional delivery")
P("Sections received a 30-minute orientation on 15–17 February 2026, using the dehydration case to introduce the "
  "arm-specific procedure. Four weekly 180-minute sessions followed between 22 February and 17 March 2026, giving 750 contact "
  "minutes and approximately 240 preparation minutes in every arm. Each session comprised a readiness check (15 minutes), "
  "individual case work (30), group work (30), plenary feedback (15), simulation prebriefing (10), two simulation-and-"
  "debriefing runs (50), consolidation (20) and summary (10). Individual case work, group work and consolidation contained "
  "the arm-specific activities (Table 1); all other segments were common to all arms.")
P("The six groups of a section completed simulation concurrently in six rooms, each staffed by one trained simulation "
  "instructor [[qualifications; state that instructors rotated across sections of all three arms]]. Each run comprised a "
  "ten-minute scenario followed by a fifteen-minute PEARLS debriefing (Eppich and Cheng, 2015). Three students had active "
  "roles in one run and two in the other, giving each student one active role per session; observers used a structured "
  "observation guide. Maps and worksheets remained outside the simulation rooms, so simulation practice was identical "
  "across arms. Feedback in every arm addressed accurate cues, age-appropriate interpretation, prioritisation, safe actions "
  "and reassessment.")

# ---------------------------------------------------------------- 2.8
H2("2.8 Standardisation, masking, fidelity and contamination")
P("Facilitators received [[hours]] of protocol training using session scripts and arm-specific facilitator guides; "
  "simulation instructors were trained in scenario operation and PEARLS debriefing. Assessment staff (standardised parents "
  "and scenario operators) were trained separately on scripts, cue timing and scripted responses. [[Assessment staff were "
  "not informed of students' allocation; confirm.]]")
P("Students and facilitators knew the learning activity. Outcome raters were masked to strategy and occasion. An "
  "independent custodian retained the code key; recordings received neutral identifiers, date overlays were removed and "
  "rating order was randomised across strategies and occasions. Students were asked not to mention their learning activity "
  "during assessment, and raters documented any disclosure. [[The statistician analysed coded strategy labels until the "
  "analysis plan was finalised; confirm.]]")
P("Printed materials were distributed and collected in each session, no electronic copies were issued, and all-arm access "
  "was provided on 9 April 2026 after follow-up. Attendance at each session, preparation engagement [[source, for example "
  "learning-management-system access and self-check quiz completion]], active simulation roles and product submission were "
  "recorded. Completion of instruction was defined as attendance at three or more of the four teaching sessions. At the "
  "follow-up visit, students reported cross-arm material use, outside AI use for this unit, additional tutoring and "
  "advance knowledge of assessment cases. These single items were analysed individually as implementation information and "
  "were not summed into a scale.")
P("All orientation and teaching sessions were audio-recorded. Eight sessions per arm (24 of 75) were [[selected by "
  "stratified random sampling across facilitators and session numbers]] and coded by independent coders using a fidelity "
  "checklist covering segment timing, delivery of the common feedback criteria, required arm components and the absence of "
  "cross-arm content. [[Number]] sessions were double-coded, with agreement reported as percentage agreement and Cohen's "
  "kappa. Deviations were logged with date, section, nature and corrective action, and fidelity is reported by arm.")

# ---------------------------------------------------------------- 2.9
H2("2.9 Outcomes and assessment procedures")
P("The primary outcome was clinical reasoning performance, measured as the post-instruction PN-CRER occasion score. "
  "Secondary outcomes were knowledge (PNNKT) and procedural accuracy (PNPSC). Table 2 gives each measure, its score range, "
  "administration and timing. Baseline simulation and skills assessment occurred on 1–10 February 2026, followed by the "
  "knowledge test on 11 February. Post-instruction simulation and skills assessment occurred on 29 March–7 April, followed "
  "by the knowledge test on 8 April. Performance assessment took place 12–23 days after each student's final teaching "
  "session. There was no delayed retention assessment.")

H3("2.9.1 Simulation-based assessment")
P("At each occasion, students individually completed one young-child case (aged 9–18 months) and one school-age case "
  "(aged 5–8 years) from the eight-case assessment bank. Case assignment and order were randomised under constraints that "
  "prevented a student from meeting the same case twice and balanced case use within sections. The expert panel reviewed "
  "the cases for comparable severity, cue timing, decision points and critical actions. Each ten-minute scenario used a "
  "standardised parent, a simulated monitor and scripted clinician responses. Three deterioration triggers occurred at "
  "fixed times, and responses to specified student actions followed the scenario script.")
P("[[Equipment, from the laboratory inventory: state the simulator used for the young-child cases and its represented "
  "age, and the simulator used for the school-age cases. State how findings that the equipment could not physically "
  "reproduce for a 9–18-month-old, including fontanelle findings, were presented (for example scripted standardised-parent "
  "report, cue card or monitor display), when they were presented and how the case key credited the corresponding "
  "assessment. See the alternative wordings in the verification log.]] Equipment and cue presentation were identical for "
  "all arms.")
P("Each case was followed by a structured explanation of up to five minutes, prompted by six prerecorded neutral prompts "
  "in a fixed order, with English audio and Arabic text. Students responded in Arabic or English, and bilingual raters "
  "judged clinical content rather than language fluency. No maps, AI, notes or feedback were available. The procedural "
  "station followed the two cases during the same visit, and knowledge testing followed completion of all performance "
  "assessments to avoid cueing case responses. Assessments were video-recorded and did not contribute to course grades. "
  "Students agreed not to discuss assessment content [[confirm undertaking]].")

H3("2.9.2 Pediatric Neurotrauma Clinical Reasoning Evaluation Rubric")
P("The investigator-developed PN-CRER contained eight domains corresponding to the Clinical Reasoning Cycle. Each domain "
  "was scored 0–3 against behavioural descriptors and case-specific keys: 0, no evidence; 1, limited or substantially "
  "incorrect evidence; 2, substantially correct but incomplete evidence; and 3, complete and accurate evidence meeting the "
  "keyed criteria. The scoring manual specified the evidence source for each domain: [[observed performance, the structured "
  "explanation or both; state the source for each of the eight domains exactly as in the final manual]]. Case keys "
  "distinguished performed from stated elements. [[A performable action that was stated but not performed could not score "
  "above ___ in the actions domain, and a keyed critical safety error capped the affected domain at ___; reproduce the "
  "caps applied to the recordings.]] Domain scores summed to a case total of 0–24, and the mean of the two case totals "
  "formed the occasion score (0–24); higher scores indicated better assessed performance. PN-CRER was a new instrument; "
  "measurement properties published for the Clinical Reasoning Evaluation Simulation Tool were not attributed to it.")

H3("2.9.3 Pediatric Neurotrauma Nursing Knowledge Test")
P("The investigator-developed PNNKT comprised two 40-item single-best-answer forms with [[number]] options per item. The "
  "forms were assembled from a piloted 96-item bank to a common blueprint matching curriculum content, cognitive level and "
  "item difficulty, following established item-writing principles (Haladyna et al., 2002). Students were randomly "
  "assigned to an A–B or B–A sequence, so no item was repeated for any student. Raw scores ranged from 0 to 40. Testing was "
  "proctored, computer-based and simultaneous in two adjacent halls, with a sixty-minute limit, and results were withheld "
  "until follow-up was complete.")
P("Random assignment of forms at baseline created a random-groups equating design (Kolen and Brennan, 2014). [[Form A was "
  "the reference form. Form B scores were placed on the Form A scale by linear equating, eA(x) = μA + (σA/σB)(x − μB), with "
  "means and standard deviations estimated from the baseline random groups (Form A, n = ___; Form B, n = ___). The resulting "
  "transformation, equated B = ___ × raw B + ___, was applied unchanged to follow-up Form B scores, and Form A scores were not "
  "transformed. Equated scores outside 0–40 were ___ (retained / truncated).]] All knowledge analyses used scores on the "
  "reference-form metric, and the sensitivity analyses included test-form sequence as a covariate. KR-20 was calculated for "
  "each form at baseline.")

H3("2.9.4 Pediatric Neurotrauma Procedural Skills Checklist")
P("The investigator-developed PNPSC contained fifteen items [[list the tasks in Table 2 or the materials]] scored 0 for "
  "absent performance, 1 for incomplete or incorrect performance and 2 for correct performance, giving a total of 0–30. The "
  "ten-minute station used explicit task instructions and the same tasks on both occasions, so it assessed execution rather "
  "than selection of care priorities. Higher scores indicated more accurate procedural performance. Recordings were rated "
  "by masked raters.")

# ---------------------------------------------------------------- 2.10
H2("2.10 Instrument development and measurement evidence")
P("Instrument development followed an argument-based approach to validity. Evidence was organised by the scoring, "
  "generalisation and extrapolation inferences that link observed performance to the intended score interpretation "
  "(Kane, 2013; Cook et al., 2015). Table 3 summarises the assumptions tested and the evidence obtained. All results are "
  "interpreted as local evidence for group comparisons in this cohort, not as general proof of instrument validity.")

H3("2.10.1 Content review and translation")
P("Eight independent experts reviewed relevance, clarity, coverage, age appropriateness and assessment-case comparability "
  "in two rounds between 16 December 2025 and 8 January 2026. The panel comprised three pediatric nursing academics, two "
  "pediatric critical care nurse specialists, a pediatric neurosurgeon, a pediatric emergency physician and a nursing "
  "measurement specialist. Relevance was rated on a four-point scale, with ratings of 3 or 4 counted as relevant. The "
  "item-level content validity index (I-CVI) was the proportion of experts rating an element relevant. Scale-level indices "
  "were calculated by the averaging method (S-CVI/Ave) and the universal-agreement method (S-CVI/UA), and a modified kappa "
  "adjusted each I-CVI for chance agreement (Lynn, 1986; Polit et al., 2007). With eight experts, the acceptability "
  "criteria were an I-CVI of at least 0.78 and an S-CVI/Ave of at least 0.90. Elements below these criteria or with "
  "substantive comments were revised and re-rated in round two. All indices were computed directly from the original expert "
  "rating forms. Instructions and bilingual prompts underwent forward translation, back-translation, expert reconciliation "
  "and cognitive testing with [[number]] students (Beaton et al., 2000).")

H3("2.10.2 Piloting")
P("External piloting on 18–22 January 2026 involved twenty-four fourth-year nursing students and eight nurses for "
  "performance assessment, 104 fourth-year volunteers for the knowledge bank and ten volunteers per arm for a teaching "
  "rehearsal. Pilot participants were not part of the study cohort and [[signed confidentiality undertakings and did not "
  "retain materials]]. Performance piloting examined cue delivery, timing, comprehension, scoring interpretation and score "
  "distributions. Knowledge items were examined for difficulty, discrimination and distractor function, with retention "
  "criteria of [[difficulty 0.30–0.80 and corrected point-biserial correlation of at least 0.20; confirm]], and KR-20 was "
  "calculated for each assembled form.")

H3("2.10.3 Raters, rating design and agreement")
P("Four bilingual pediatric nursing raters [[qualifications]], uninvolved in teaching and instrument development, "
  "completed rater training on 12–13 April 2026 using pilot benchmark recordings. They reached [[calibration criterion]] "
  "before live rating, which took place from [[start date]] to [[end date]]. Three primary raters scored recordings "
  "allocated at random [[within arm and occasion strata]], and primary-rater scores were used in the primary analysis. A "
  "fourth rater independently scored 360 case recordings selected at random [[within arm and occasion strata]]. All four "
  "raters scored a common subset of 120 recordings from sixty students; these recordings were fully crossed with raters, "
  "but students were incompletely crossed with the eight cases. For procedural performance, 225 recordings were "
  "independently double-rated [[state rater pairing]]. Benchmark recordings were inserted [[frequency]] to monitor drift, "
  "and recalibration was documented.")
P("Agreement statistics were matched to the rating design and score unit (Koo and Li, 2016). Because the second rating in "
  "the 360-recording sample was paired with different primary raters, agreement there was estimated with a one-way "
  "random-effects, absolute-agreement, single-rater intraclass correlation on case totals. The common subset, in which all "
  "four raters scored every recording, used a two-way random-effects, absolute-agreement, single-rater intraclass "
  "correlation. Procedural agreement used the model corresponding to its rater pairing. Students contributed more than one "
  "recording, so 95% confidence intervals were obtained by cluster bootstrap resampling of students (2,000 resamples). "
  "Domain-level agreement was summarised as exact and within-one-point agreement.")
P("A generalisability study on the common subset estimated variance components by restricted maximum likelihood in a "
  "crossed random-effects model for students (p), cases (c), raters (r), student-by-case (pc), student-by-rater (pr), "
  "case-by-rater (cr) and residual (pcr,e) effects (Briesch et al., 2014). This approach accommodates the unbalanced "
  "incomplete crossing of students with cases. The relative coefficient for n′c cases and n′r raters was "
  "σ²(p) / [σ²(p) + σ²(pc)/n′c + σ²(pr)/n′r + σ²(pcr,e)/(n′c n′r)]. The dependability coefficient additionally included "
  "σ²(c)/n′c, σ²(r)/n′r and σ²(cr)/(n′c n′r) in the error term. Coefficients were calculated for the operational design of "
  "two cases scored by one rater and reported with bootstrap confidence intervals.")
P("For a separate convergent comparison, the Lasater Clinical Judgment Rubric (LCJR) was used with the author's "
  "permission (Lasater, 2007). The LCJR has eleven dimensions rated at four levels, giving totals of 11–44. Its primary "
  "rater [[identify; state that this rater did not score PN-CRER on the same recordings]] scored 270 case recordings "
  "[[selected at random within arm and occasion strata]], and a second rater independently scored 135 of these recordings "
  "to provide LCJR agreement information. The association between PN-CRER and LCJR case totals was estimated with "
  "students as the sampling unit, using [[a repeated-measures correlation / a multilevel bivariate model]] with cluster "
  "bootstrap confidence intervals. [[If an expected direction and minimum magnitude were stated in the protocol, give them "
  "here.]]")

# ---------------------------------------------------------------- 2.11
H2("2.11 Statistical analysis")
H3("2.11.1 Estimands and contrasts")
P("The primary estimand was the adjusted difference in mean post-instruction PN-CRER occasion score between allocated "
  "strategies among first-time-enrolled third-year students in this cohort. It was defined regardless of attendance, "
  "cross-arm material use or outside AI use, which were handled by a treatment-policy strategy (International Council for "
  "Harmonisation, 2019). The principal contrast compared AI-generated with student-generated maps at a two-sided "
  "significance level of 0.05. The two comparisons with worksheets were supporting contrasts and were adjusted together "
  "using a multivariate-t, Dunnett-type procedure (Hothorn et al., 2008). Results are expressed as adjusted mean "
  "differences with 95% confidence intervals. Within-arm change tests were not used to establish between-arm differences.")

H3("2.11.2 Primary analysis")
P("The primary linear mixed model analysed the primary raters' post-instruction case totals, with two observations per "
  "student. Because case assignment was balanced, the strategy contrasts are on the 0–24 occasion-score metric. Fixed "
  "effects were strategy; facilitator, the allocation stratum; baseline reasoning, defined as the mean of the student's two "
  "baseline case totals; GPA; baseline knowledge; English course grade; previous simulation sessions; previous mapping use; "
  "previous AI use; case age group; case order; and rater. The four allocation-balancing variables were therefore included "
  "as adjustment variables. Random intercepts represented section, group within section, student and assessment case, "
  "with case crossed with student. Design-based random effects for section and group were retained irrespective of their "
  "estimated magnitude. Estimation used restricted maximum likelihood, with degrees of freedom approximated by the "
  "Kenward–Roger method, which is appropriate for analyses with few clusters (Kenward and Roger, 1997; Leyrat et al., "
  "2018).")
P("To aid interpretation, adjusted differences were also divided by the total standard deviation of post-instruction "
  "occasion scores estimated from an unconditional multilevel model (Hedges, 2007). Section and group intraclass "
  "correlations from the same model are reported to inform future studies.")

H3("2.11.3 Secondary outcomes")
P("Knowledge and procedural scores were analysed in student-level linear mixed models with random intercepts for section "
  "and group within section. Each model adjusted for the corresponding baseline score, facilitator and the prognostic "
  "variables listed above; the knowledge model also included test-form sequence, and the procedural model included rater. "
  "Secondary comparisons were interpreted as supporting estimates without multiplicity adjustment.")

H3("2.11.4 Missing data")
P("Students remained classified by their allocated strategy irrespective of attendance or contamination. Baseline outcome "
  "assessments were complete, and [[baseline covariates were complete; confirm]]. The primary likelihood analysis used all "
  "observed post-instruction outcomes and is valid under a missing-at-random assumption conditional on the variables in the "
  "model. Students with no post-instruction assessment contributed no follow-up observation, and missingness and its "
  "reasons are reported by arm (Figure 1). In a sensitivity analysis, multilevel multiple imputation [[number of imputed "
  "data sets; method or package]] was performed separately by arm and included the analysis-model variables, baseline "
  "outcomes, attendance and preparation engagement. Estimates were combined by Rubin's rules. Departures from missing at "
  "random were then examined by delta adjustment, subtracting 0 to 0.5 standard deviations from imputed values in each arm "
  "separately (Cro et al., 2020). [[State how the two students who withdrew were handled, consistent with their consent.]]")

H3("2.11.5 Sensitivity analyses")
P("Sensitivity analyses examined (i) section-level estimates; (ii) alternative rater aggregation, averaging both ratings "
  "where a recording was double-rated; (iii) additional adjustment for timetable position and teaching-to-assessment "
  "interval [[confirm that these were the timing variables]] and for clinical exposure to children with head injury "
  "between assessments; (iv) restriction to students who completed instruction; and (v) restriction to students "
  "reporting no cross-arm material use or outside AI use. Restricted analyses were interpreted as comparisons of selected "
  "subgroups rather than as estimates of the primary estimand. Analyses identified in version 2 of the plan are labelled "
  "as such in the Results.")

H3("2.11.6 Model checking, descriptive analysis and software")
P("Baseline characteristics are summarised by arm with standardised mean differences; significance tests were not used "
  "for baseline comparisons. Residual distributions, heterogeneity of residual variance by arm, convergence and singular "
  "fits were examined, and [[state the prespecified response, for example refitting with arm-specific residual variances]]. "
  "Analyses used R version 4.3.1 [[list the packages used, for example lme4, pbkrtest, emmeans, multcomp and the imputation "
  "package]] (Bates et al., 2015). Estimates are interpreted as adjusted comparisons of instructional strategies within this "
  "cohort; nonrandom allocation and five sections per arm limit causal inference.")

# ---------------------------------------------------------------- 2.12
H2("2.12 Ethics")
P("The Ethics Committee of the Faculty of Nursing, Tanta University, approved the study on 15 December 2025 "
  "(EC-2025-089), [[and the study was conducted in accordance with the Declaration of Helsinki]]. Written consent covered "
  "research data and assessment recording, and optional future use of recordings for rater training was consented "
  "separately. An independent researcher held consent records, and teaching staff did not know who had consented. "
  "Participation had no effect on grades. Students could stop an assessment if distressed and were offered support "
  "through [[service]]. All arms received all teaching materials after follow-up.")

H2("2.13 Data management and availability of materials")
P("Coded data and recordings were stored on encrypted, access-controlled institutional systems, with the identity key held "
  "separately by the independent custodian. Data entry was [[double-entered or source-verified]], and the database was "
  "locked on [[date]]. Identifiable recordings were not entered into any AI system or included in data sharing, and data "
  "will be retained for [[period]]. The protocol, analysis-plan versions, allocation and analysis code, case scripts, "
  "reviewed maps and their change logs, generation prompts, original instrument forms and scoring keys are available from "
  "the corresponding author on reasonable request, subject to institutional requirements. The PN-CRER, PNNKT and PNPSC are "
  "investigator-developed instruments owned by [[owner]], and they are shared under [[terms]].")

# ---------------------------------------------------------------- Tables
H1("Tables")
H2("Table 1 Comparison of the three instructional strategies (TIDieR)")
W4 = [1700, 2553, 2553, 2554]
table(doc, ["Feature", "AI-generated maps", "Student-generated maps", "Worksheets without maps"], [
    ["Starting material", "Clinically reviewed AI-generated map of the index case", "Central topic and eight main branches; no completed content", "Eight linear case-analysis headings"],
    ["Individual case work (30 min)", "Examine and annotate the map; justify priorities", "Select information and construct a scaffolded map", "Analyse information and justify priorities in writing"],
    ["Group work (30 min)", "Revise the map and record reasons for each revision", "Compare and consolidate maps; record revisions", "Consolidate written analysis; record reasons"],
    ["Clinical relationships", "Add and explain cue–problem–action links", "Construct and explain cue–problem–action links", "Explain relationships in text"],
    ["Consolidation (20 min)", "[[Describe]]", "[[Describe]]", "[[Describe]]"],
    ["Product collected", "Annotated and revised map", "Group-consolidated map", "Consolidated worksheet"],
    ["Direct AI interaction", "None", "None", "None"],
], W4)
para(doc, "Common to all arms: flipped preparation (about 60 minutes per session), orientation (30 minutes), four 180-minute "
     "sessions, readiness check, plenary feedback using common criteria, simulation prebriefing, two simulation runs with "
     "PEARLS debriefing, summary, facilitators and simulation instructors. Materials were collected after each session and "
     "released to all arms on 9 April 2026. No maps, AI or notes were available during assessment. The comparison concerned "
     "representation source and learner activity together.", size=9.5)

H2("Table 2 Outcomes, instruments and assessment schedule")
W5 = [1350, 2500, 2310, 1600, 1600]
table(doc, ["Outcome", "Instrument and score", "Administration and rating", "Timing", "Analysis"], [
    ["Primary: clinical reasoning", "PN-CRER; eight domains scored 0–3 against case keys; case total 0–24; occasion score = mean of two cases (0–24)",
     "Individual simulation (young-child and school-age case) with structured explanation; masked video rating; no instructional aids",
     "1–10 February and 29 March–7 April 2026", "Case-level linear mixed model; adjusted mean difference"],
    ["Secondary: knowledge", "PNNKT; two 40-item forms; raw 0–40; equated to reference-form metric",
     "Proctored, 60-minute computer test after performance assessments; random form sequence",
     "11 February and 8 April 2026", "Student-level linear mixed model"],
    ["Secondary: procedural accuracy", "PNPSC; fifteen items scored 0–2; total 0–30 [[list tasks]]",
     "Ten-minute instructed-task station; masked video rating", "Same visit as the reasoning cases", "Student-level linear mixed model"],
    ["Implementation", "Attendance; preparation engagement; active simulation roles; product submission; fidelity checklist; single-item contamination reports",
     "Registers, [[system logs]], audio-recorded sessions coded by independent coders", "Throughout teaching; contamination items at follow-up", "Descriptive by arm"],
], W5)
para(doc, "PN-CRER, Pediatric Neurotrauma Clinical Reasoning Evaluation Rubric; PNNKT, Pediatric Neurotrauma Nursing "
     "Knowledge Test; PNPSC, Pediatric Neurotrauma Procedural Skills Checklist. No delayed retention assessment was "
     "conducted.", size=9.5)

H2("Table 3 Measurement evidence organised by validity inference")
W4b = [1500, 2700, 2860, 2300]
table(doc, ["Inference", "Assumption", "Evidence collected", "Analysis"], [
    ["Scoring", "Domains, descriptors and case keys represent the reasoning cycle accurately and can be applied consistently",
     "Two-round expert review (n = 8); scoring manual with evidence sources and caps; rater training and benchmark calibration",
     "I-CVI, S-CVI/Ave, S-CVI/UA, modified kappa; calibration records"],
    ["Scoring", "Bilingual prompts and explanations convey the same meaning in Arabic and English",
     "Forward and back-translation, reconciliation and cognitive testing", "Documented revisions"],
    ["Generalisation", "Scores generalise across raters", "360 recordings double-rated; 120 recordings rated by all four raters; 225 procedural recordings double-rated",
     "Design-matched ICCs with cluster-bootstrap CIs; exact and adjacent agreement"],
    ["Generalisation", "Scores generalise across cases and raters for the two-case occasion score",
     "Common subset of 120 recordings, students incompletely crossed with cases", "REML variance components; relative and dependability coefficients"],
    ["Generalisation", "Knowledge forms are reliable and comparable", "Pilot item analysis (n = 104); baseline random-groups data",
     "Item statistics; KR-20 by form; linear equating"],
    ["Extrapolation", "PN-CRER scores relate to an established clinical judgement measure", "LCJR on 270 recordings; 135 double-rated",
     "Student-level association with cluster-bootstrap CIs; LCJR agreement"],
], W4b)
para(doc, "CI, confidence interval; ICC, intraclass correlation coefficient; I-CVI, item-level content validity index; "
     "LCJR, Lasater Clinical Judgment Rubric; REML, restricted maximum likelihood; S-CVI/Ave and S-CVI/UA, scale-level "
     "content validity index by averaging and universal agreement.", size=9.5)

H2("Table 4 Sequence of study procedures")
W2 = [2700, 6660]
table(doc, ["Date", "Procedure"], [
    ["15 November 2025", "Protocol and statistical analysis plan, version 1"],
    ["15 December 2025", "Ethics approval (EC-2025-089)"],
    ["16 December 2025–8 January 2026", "Two-round expert review of instruments and assessment cases"],
    ["5–8 January 2026", "Generation of candidate AI outlines"],
    ["11–15 January 2026", "Clinical review of selected maps"],
    ["18–22 January 2026", "External piloting"],
    ["19 January 2026", "Map content frozen"],
    ["[[Date]]", "Recruitment and written consent"],
    ["1–10 February 2026", "Baseline simulation and procedural assessment"],
    ["11 February 2026", "Baseline knowledge test"],
    ["12 February 2026", "Nonrandom section allocation"],
    ["15–17 February 2026", "Orientation"],
    ["22 February–17 March 2026", "Four weekly teaching sessions"],
    ["29 March–7 April 2026", "Post-instruction simulation and procedural assessment"],
    ["8 April 2026", "Post-instruction knowledge test"],
    ["9 April 2026", "All teaching materials released to all arms"],
    ["12–13 April 2026", "Rater training"],
    ["[[Dates]]", "Masked rating of baseline and post-instruction recordings"],
    ["[[Date and time]]", "Rating completed; database locked"],
    ["30 June 2026", "Analysis plan version 2 [[and release of the strategy code key; give the order]]"],
], W2)

# ---------------------------------------------------------------- Figure
H1("Figure")
H2("Figure 1 Participant flow through the three-arm quasi-experimental study")
doc.add_picture(FIG, width=Inches(6.3))
para(doc, "All fifteen sections remained in the study. Allocation was nonrandom. The attendance threshold described completion "
     "of instruction and did not exclude students from the allocated-strategy analysis. Twenty-one students lacked "
     "post-instruction assessment: illness (8), holiday travel (7), clinical-placement conflict (4) and withdrawal (2). The "
     "429 students with follow-up supplied the observed post-instruction outcomes; all 450 allocated students were included "
     "in the imputation sensitivity analysis. [[Complete the highlighted counts by arm from the study records.]]", size=9.5)

# ---------------------------------------------------------------- Declaration
H1("Declaration of generative AI use in manuscript preparation")
P("[[During preparation of this manuscript, the authors used ChatGPT (OpenAI) and Claude (Anthropic) to assist with "
  "organisation, language revision and checking of methodological reporting. The authors reviewed and edited all material, "
  "verified the study records and references, and take full responsibility for the content of the publication. Use this "
  "statement only after that review has occurred, and list only the tools actually used.]] This declaration is separate "
  "from the GPT-4o use described in Section 2.6, which formed part of the intervention.")

# ---------------------------------------------------------------- References
H1("References")
REFS = [
 "Akutay, S., Yüceler Kaçmaz, H., Kahraman, H., 2024. The effect of artificial intelligence supported case analysis on nursing students' case management performance and satisfaction: A randomized controlled trial. Nurse Education in Practice 80, 104142. https://doi.org/10.1016/j.nepr.2024.104142",
 "Arkan, B., Erbay Dallı, Ö., Varol, B., 2025. The impact of ChatGPT training in the nursing process on nursing students' problem-solving skills, attitudes towards artificial intelligence, competency, and satisfaction levels: Single-blind randomized controlled study. Nurse Education Today 152, 106765. https://doi.org/10.1016/j.nedt.2025.106765",
 "Bates, D., Mächler, M., Bolker, B., Walker, S., 2015. Fitting linear mixed-effects models using lme4. Journal of Statistical Software 67(1), 1–48. https://doi.org/10.18637/jss.v067.i01",
 "Bayzat, C., Dinc, L., 2025. Turkish adaptation of the clinical reasoning evaluation simulation tool (CREST-TR): a validity and reliability study. BMC Nursing 24, 819. https://doi.org/10.1186/s12912-025-03503-0",
 "Beaton, D.E., Bombardier, C., Guillemin, F., Ferraz, M.B., 2000. Guidelines for the process of cross-cultural adaptation of self-report measures. Spine 25, 3186–3191. https://doi.org/10.1097/00007632-200012150-00014",
 "Briesch, A.M., Swaminathan, H., Welsh, M., Chafouleas, S.M., 2014. Generalizability theory: a practical guide to study design, implementation, and interpretation. Journal of School Psychology 52(1), 13–35. https://doi.org/10.1016/j.jsp.2013.11.008",
 "Chan, K., Farrell, C.A., Chauvin-Kimoff, L., 2025. Management of the paediatric patient with acute head trauma. Paediatrics & Child Health 30, 641–647. https://doi.org/10.1093/pch/pxaf032",
 "Chang, C.-Y., Su, W.-S., 2025. The effect of a generative AI-based teaching strategy on building students' competency. Journal of Nursing Education 64(6), 346–355. https://doi.org/10.3928/01484834-20250129-05",
 "Chen, T.-Y., Hung, C.-C., 2025. An integrated self-regulated learning and flipped classroom approach for teaching nursing skills to undergraduate nursing students: A randomized controlled study. Nurse Education in Practice 87, 104445. https://doi.org/10.1016/j.nepr.2025.104445",
 "Cheng, A., Kessler, D., Mackinnon, R., Chang, T.P., Nadkarni, V.M., Hunt, E.A., Duval-Arnould, J., Lin, Y., Cook, D.A., Pusic, M., Hui, J., Moher, D., Egger, M., Auerbach, M., 2016. Reporting guidelines for health care simulation research: extensions to the CONSORT and STROBE statements. Advances in Simulation 1, 25. https://doi.org/10.1186/s41077-016-0025-y",
 "Chiannilkulchai, N., Thongpo, P., 2026. Effect of online flipped classroom and case-based learning approaches on nursing students' training in primary survey for neurotrauma care: a quasi-experimental study. BMC Medical Education. https://doi.org/10.1186/s12909-026-10226-6",
 "Cook, D.A., Brydges, R., Ginsburg, S., Hatala, R., 2015. A contemporary approach to validity arguments: a practical guide to Kane's framework. Medical Education 49(6), 560–575. https://doi.org/10.1111/medu.12678",
 "Cro, S., Morris, T.P., Kenward, M.G., Carpenter, J.R., 2020. Sensitivity analysis for clinical trials with missing continuous outcome data using controlled multiple imputation: A practical guide. Statistics in Medicine 39(21), 2815–2842. https://doi.org/10.1002/sim.8569",
 "Des Jarlais, D.C., Lyles, C., Crepaz, N., TREND Group, 2004. Improving the reporting quality of nonrandomized evaluations of behavioral and public health interventions: the TREND statement. American Journal of Public Health 94, 361–366. https://doi.org/10.2105/AJPH.94.3.361",
 "Dwiarini, M., Hwang, G.-J., Chang, C.-Y., 2026. Effects of a generative AI-assisted ADPIE Framework on critical thinking, learning self-efficacy, and clinical reasoning in nursing education: a quasi-experimental study. BMC Nursing 25, 779. https://doi.org/10.1186/s12912-026-04840-4",
 "Eppich, W., Cheng, A., 2015. Promoting Excellence and Reflective Learning in Simulation (PEARLS): development and rationale for a blended approach to health care simulation debriefing. Simulation in Healthcare 10, 106–115. https://doi.org/10.1097/SIH.0000000000000072",
 "Fan, Y., Tang, L., Le, H., Shen, K., Tan, S., Zhao, Y., Shen, Y., Li, X., Gašević, D., 2025. Beware of metacognitive laziness: Effects of generative artificial intelligence on learning motivation, processes, and performance. British Journal of Educational Technology 56(2), 489–530. https://doi.org/10.1111/bjet.13544",
 "Fiorella, L., Mayer, R.E., 2016. Eight ways to promote generative learning. Educational Psychology Review 28(4), 717–741. https://doi.org/10.1007/s10648-015-9348-9",
 "Haladyna, T.M., Downing, S.M., Rodriguez, M.C., 2002. A review of multiple-choice item-writing guidelines for classroom assessment. Applied Measurement in Education 15, 309–333. https://doi.org/10.1207/S15324818AME1503_5",
 "Hedges, L.V., 2007. Effect sizes in cluster-randomized designs. Journal of Educational and Behavioral Statistics 32(4), 341–370. https://doi.org/10.3102/1076998606298043",
 "Hoffmann, T.C., Glasziou, P.P., Boutron, I., et al., 2014. Better reporting of interventions: template for intervention description and replication (TIDieR) checklist and guide. BMJ 348, g1687. https://doi.org/10.1136/bmj.g1687",
 "Hothorn, T., Bretz, F., Westfall, P., 2008. Simultaneous inference in general parametric models. Biometrical Journal 50(3), 346–363. https://doi.org/10.1002/bimj.200810425",
 "Ibrahim, R.K., Hendy, A., 2026. The role of digital mind maps in boosting creativity and critical thinking among nursing students: a quasi-experimental study. Teaching and Learning in Nursing 21(2), e525–e533. https://doi.org/10.1016/j.teln.2025.10.026",
 "International Council for Harmonisation of Technical Requirements for Pharmaceuticals for Human Use, 2019. Addendum on estimands and sensitivity analysis in clinical trials to the guideline on statistical principles for clinical trials, E9(R1). ICH, Geneva. https://database.ich.org/sites/default/files/E9-R1_Step4_Guideline_2019_1203.pdf",
 "Kane, M.T., 2013. Validating the interpretations and uses of test scores. Journal of Educational Measurement 50, 1–73. https://doi.org/10.1111/jedm.12000",
 "Kenward, M.G., Roger, J.H., 1997. Small sample inference for fixed effects from restricted maximum likelihood. Biometrics 53, 983–997. https://doi.org/10.2307/2533558",
 "Khanna, S.K., Kumar, A., Katiyar, A.K., Mishra, K., 2025. Clinical profile, management, and outcome of pediatric neurotrauma: a multicentric observational study. Journal of Trauma and Injury 38, 22–31. https://doi.org/10.20408/jti.2024.0080",
 "Kochanek, P.M., Tasker, R.C., Carney, N., et al., 2019. Guidelines for the Management of Pediatric Severe Traumatic Brain Injury, Third Edition: Update of the Brain Trauma Foundation Guidelines. Pediatric Critical Care Medicine 20(3S Suppl 1), S1–S82. https://doi.org/10.1097/PCC.0000000000001735",
 "Kolen, M.J., Brennan, R.L., 2014. Test Equating, Scaling, and Linking: Methods and Practices, third ed. Springer, New York. https://doi.org/10.1007/978-1-4939-0317-7",
 "Koo, T.K., Li, M.Y., 2016. A guideline of selecting and reporting intraclass correlation coefficients for reliability research. Journal of Chiropractic Medicine 15, 155–163. https://doi.org/10.1016/j.jcm.2016.02.012",
 "Lasater, K., 2007. Clinical judgment development: using simulation to create an assessment rubric. Journal of Nursing Education 46(11), 496–503. https://doi.org/10.3928/01484834-20071101-04",
 "Lenski, S., Elsner, S., Großschedl, J., 2022. Comparing construction and study of concept maps – An intervention study on learning outcome, self-evaluation and enjoyment through training and learning. Frontiers in Education 7, 892312. https://doi.org/10.3389/feduc.2022.892312",
 "Levett-Jones, T., Hoffman, K., Dempsey, J., Jeong, S.Y.-S., Noble, D., Norton, C.A., Roche, J., Hickey, N., 2010. The 'five rights' of clinical reasoning: an educational model to enhance nursing students' ability to identify and manage clinically 'at risk' patients. Nurse Education Today 30, 515–520. https://doi.org/10.1016/j.nedt.2009.10.020",
 "Leyrat, C., Morgan, K.E., Leurent, B., Kahan, B.C., 2018. Cluster randomized trials with a small number of clusters: which analyses should be used?. International Journal of Epidemiology 47, 321–331. https://doi.org/10.1093/ije/dyx169",
 "Liaw, S.Y., Rashasegaran, A., Wong, L.F., Deneen, C.C., Cooper, S., Levett-Jones, T., Goh, H.S., Ignacio, J., 2018. Development and psychometric testing of a Clinical Reasoning Evaluation Simulation Tool (CREST) for assessing nursing students' abilities to recognize and respond to clinical deterioration. Nurse Education Today 62, 74–79. https://doi.org/10.1016/j.nedt.2017.12.009",
 "Liaw, S.Y., Scherpbier, A., Rethans, J.-J., Klainin-Yobas, P., 2012. Assessment for simulation learning outcomes: a comparison of knowledge and self-reported confidence with observed clinical performance. Nurse Education Today 32, e35–e39. https://doi.org/10.1016/j.nedt.2011.10.006",
 "Lynn, M.R., 1986. Determination and quantification of content validity. Nursing Research 35(6), 382–385. https://doi.org/10.1097/00006199-198611000-00017",
 "Mesk, Z., Qtait, M., Alqaissi, N., Amro, A., Asafrah, M., Zeedat, F., Zeedat, H., Daoud, M., Asafrah, D., Hmedat, N., 2026. Effectiveness of mind mapping in undergraduate nursing education: A quasi-experimental study. SAGE Open Medicine 14, 20503121261462854. https://doi.org/10.1177/20503121261462854",
 "National Institute for Health and Care Excellence, 2023. Head injury: assessment and early management. NICE guideline NG232. https://www.nice.org.uk/guidance/ng232",
 "Phillips, A.C., Lewis, L.K., McEvoy, M.P., Galipeau, J., Glasziou, P., Moher, D., Tilson, J.K., Williams, M.T., 2016. Development and validation of the guideline for reporting evidence-based practice educational interventions and teaching (GREET). BMC Medical Education 16, 237. https://doi.org/10.1186/s12909-016-0759-1",
 "Polit, D.F., Beck, C.T., Owen, S.V., 2007. Is the CVI an acceptable indicator of content validity? Appraisal and recommendations. Research in Nursing & Health 30, 459–467. https://doi.org/10.1002/nur.20199",
 "Shin, H., De Gagne, J.C., Kim, S.S., Hong, M., 2024. The impact of artificial intelligence-assisted learning on nursing students' ethical decision-making and clinical reasoning in pediatric care: A quasi-experimental study. CIN: Computers, Informatics, Nursing 42(10), 704–711. https://doi.org/10.1097/CIN.0000000000001177",
 "Sweller, J., van Merriënboer, J.J.G., Paas, F., 2019. Cognitive architecture and instructional design: 20 years later. Educational Psychology Review 31, 261–292. https://doi.org/10.1007/s10648-019-09465-5",
 "Wang, L., Wang, Y., Wang, X., Xue, C., 2023. Effects of mind mapping based on standardized patient program in patient education among postgraduate nursing students in clinical setting. BMC Medical Education 23, 982. https://doi.org/10.1186/s12909-023-04944-4",
 "Watts, P.I., et al., 2021. Healthcare Simulation Standards of Best Practice: Simulation design. Clinical Simulation in Nursing 58, 14–21. https://doi.org/10.1016/j.ecns.2021.08.009",
 "Yang, K.-H., Chen, H., Liu, C.-J., Zhang, F.-F., Jiang, X.-L., 2022. Effects of reflective learning based on visual mind mapping in the fundamentals of nursing course: A quasi-experimental study. Nurse Education Today 119, 105566. https://doi.org/10.1016/j.nedt.2022.105566",
]
for r in REFS:
    p = para(doc, r)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)

doc.save(OUT)
print("saved", OUT)
