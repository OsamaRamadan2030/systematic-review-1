"""Build the PRISMA 2020 checklist and the cover letter for the Healthcare submission."""
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from docx_common import normalise, Inline, add_runs, set_cell_borders, set_cell_margins, set_cell_shading

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "submission_healthcare")
TITLE = ("Access or Learning? A Systematic Review of Speech and Braille Assistive Technologies "
         "for Children and Adolescents with Disabilities")
FONT = "Palatino Linotype"


def base_doc(landscape=False):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(2.0))
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    return doc


def simple_table(doc, header, rows, widths_cm, size=8.5, group_rows=()):
    inline = Inline({}, [])
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.autofit = False
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "8"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for j, w in enumerate(widths_cm):
        for row in t.rows:
            row.cells[j].width = Cm(w)

    def put(cell, text, bold=False):
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_runs(p, inline.parse(text, {"b": True} if bold else None), size=size)

    for j, h in enumerate(header):
        put(t.rows[0].cells[j], h, bold=True)
        set_cell_borders(t.rows[0].cells[j], bottom=4)
    for i, r in enumerate(rows):
        cells = t.rows[i + 1].cells
        if i in group_rows:
            m = cells[0].merge(cells[-1])
            put(m, r[0], bold=True)
            set_cell_shading(m, "F2F2F2")
            continue
        for j, v in enumerate(r):
            put(cells[j], v)
    return t


PRISMA = [
    ["TITLE"],
    ["Title", "1", "Identify the report as a systematic review.", "Title page (title ends “A Systematic Review of …”; article type: Systematic Review)."],
    ["ABSTRACT"],
    ["Abstract", "2", "See the PRISMA 2020 for Abstracts checklist below.", "Abstract (structured; 250 words); Table 2 of this file."],
    ["INTRODUCTION"],
    ["Rationale", "3", "Describe the rationale for the review in the context of existing knowledge.", "Section 1, paragraphs 1–4."],
    ["Objectives", "4", "Provide an explicit statement of the objective(s) or question(s) the review addresses.", "Section 1, final paragraph."],
    ["METHODS"],
    ["Eligibility criteria", "5", "Specify the inclusion and exclusion criteria and how studies were grouped for the syntheses.", "Section 2.2; grouping in Section 2.8 and Table S4 (Panel A)."],
    ["Information sources", "6", "Specify all databases, registers, websites, organisations, reference lists and other sources searched or consulted, with the date last searched.", "Section 2.3; Table S1."],
    ["Search strategy", "7", "Present the full search strategies for all databases, registers and websites, including filters and limits.", "Table S1."],
    ["Selection process", "8", "Specify the methods used to decide eligibility, including the number of reviewers and whether they worked independently, and any automation tools.", "Section 2.4 (single reviewer; no automation tools)."],
    ["Data collection process", "9", "Specify the methods used to collect data, including the number of reviewers, independence, author contact and automation tools.", "Section 2.5 (single reviewer with verification pass)."],
    ["Data items", "10a", "List and define all outcomes for which data were sought and the methods used to decide which results to collect.", "Section 2.2 (Outcomes); Section 2.6; Section 2.7 (target result)."],
    ["", "10b", "List and define all other variables for which data were sought and describe assumptions about missing or unclear information.", "Section 2.5; Table S2."],
    ["Study risk of bias assessment", "11", "Specify the methods used to assess risk of bias, including tools, number of reviewers and independence.", "Section 2.7; Table S3."],
    ["Effect measures", "12", "Specify for each outcome the effect measure(s) used in the synthesis or presentation of results.", "Section 2.6."],
    ["Synthesis methods", "13a", "Describe the processes used to decide which studies were eligible for each synthesis.", "Section 2.8; Table S4 (Panel A)."],
    ["", "13b", "Describe any methods required to prepare the data for presentation or synthesis.", "Sections 2.5 and 2.8 (no imputation, digitisation or reconstruction)."],
    ["", "13c", "Describe any methods used to tabulate or visually display results of individual studies and syntheses.", "Section 2.8; Tables 1–2; Figures 2–4; Tables S2 and S6."],
    ["", "13d", "Describe any methods used to synthesise results and provide a rationale for the choice(s).", "Section 2.8 (SWiM; effect-direction coding; rationale for no meta-analysis)."],
    ["", "13e", "Describe any methods used to explore possible causes of heterogeneity among study results.", "Section 2.8 (descriptive exploration by learner profile, grade, comparator, exposure; access vs. learning outcomes)."],
    ["", "13f", "Describe any sensitivity analyses conducted to assess robustness of the synthesised results.", "Section 2.8 (post hoc restriction to lower-risk reports)."],
    ["Reporting bias assessment", "14", "Describe any methods used to assess risk of bias due to missing results in a synthesis.", "Section 2.9."],
    ["Certainty assessment", "15", "Describe any methods used to assess certainty (or confidence) in the body of evidence for an outcome.", "Section 2.9."],
    ["RESULTS"],
    ["Study selection", "16a", "Describe the results of the search and selection process, ideally using a flow diagram.", "Section 3.1; Figure 1."],
    ["", "16b", "Cite studies that might appear to meet the inclusion criteria but were excluded, and explain why.", "Section 3.1; Table S4 (Panels D and E)."],
    ["Study characteristics", "17", "Cite each included study and present its characteristics.", "Section 3.2; Table 1; Table S2."],
    ["Risk of bias in studies", "18", "Present assessments of risk of bias for each included study.", "Section 3.3; Figure 3; Table S3."],
    ["Results of individual studies", "19", "For all outcomes, present summary statistics for each group and an effect estimate with precision, where available.", "Sections 3.4.1–3.4.5; Figure 4; Tables S2 and S6 (precision reported where source data allowed)."],
    ["Results of syntheses", "20a", "For each synthesis, briefly summarise the characteristics and risk of bias among contributing studies.", "Sections 3.4 and 3.6; Table 2; Table S5."],
    ["", "20b", "Present results of all statistical syntheses conducted; if none, state this.", "Sections 2.8 and 3.4 (no meta-analysis; effect-direction summaries); Figure 4."],
    ["", "20c", "Present results of all investigations of possible causes of heterogeneity.", "Section 3.4.1 (grade, reading profile, comparator); Section 4.3."],
    ["", "20d", "Present results of all sensitivity analyses conducted.", "Section 3.5."],
    ["Reporting biases", "21", "Present assessments of risk of bias due to missing results for each synthesis assessed.", "Section 3.6; Table S5 (publication and non-reporting bias column)."],
    ["Certainty of evidence", "22", "Present assessments of certainty (or confidence) in the body of evidence for each outcome assessed.", "Section 3.6; Table 2; Table S5."],
    ["DISCUSSION"],
    ["Discussion", "23a", "Provide a general interpretation of the results in the context of other evidence.", "Sections 4.1–4.3."],
    ["", "23b", "Discuss any limitations of the evidence included in the review.", "Section 4.6, paragraph 2."],
    ["", "23c", "Discuss any limitations of the review processes used.", "Section 4.6, paragraph 3."],
    ["", "23d", "Discuss implications of the results for practice, policy, and future research.", "Sections 4.4–4.5; Table 3; Section 5."],
    ["OTHER INFORMATION"],
    ["Registration and protocol", "24a", "Provide registration information, including register name and number, or state that the review was not registered.", "Section 2.1 (PROSPERO CRD420261513927)."],
    ["", "24b", "Indicate where the review protocol can be accessed, or state that a protocol was not prepared.", "Section 2.1 (PROSPERO record)."],
    ["", "24c", "Describe and explain any amendments to information provided at registration or in the protocol.", "Section 2.1; Table S4 (Panel B)."],
    ["Support", "25", "Describe sources of financial or non-financial support and the role of the funders or sponsors.", "Back matter: Funding."],
    ["Competing interests", "26", "Declare any competing interests of review authors.", "Back matter: Conflicts of Interest."],
    ["Availability of data, code and other materials", "27", "Report which data, forms, code and other materials are publicly available and where.", "Back matter: Data Availability Statement; Tables S1–S6."],
]

PRISMA_A = [
    ["1", "Title", "Identify the report as a systematic review.", "Yes: “A Systematic Review of …”."],
    ["2", "Objectives", "Provide an explicit statement of the main objective(s) or question(s).", "Yes: Background/Objectives."],
    ["3", "Eligibility criteria", "Specify the inclusion and exclusion criteria.", "Yes: Methods (designs; learner-level outcomes; school-age learners with disabilities)."],
    ["4", "Information sources", "Specify the information sources and the date each was last searched.", "Yes: Methods (four databases to 23 July 2026; Google Scholar; citation searching)."],
    ["5", "Risk of bias", "Specify the methods used to assess risk of bias.", "Yes: Methods (RoB 2, ROBINS-I, WWC standards)."],
    ["6", "Synthesis of results", "Specify the methods used to present and synthesise results.", "Yes: Methods (synthesis without meta-analysis; effect-direction coding; GRADE)."],
    ["7", "Included studies", "Give the total number of included studies and participants and summarise relevant characteristics.", "Yes: Results (27 reports; ≤26 studies; ≤577 learners with disabilities)."],
    ["8", "Synthesis of results", "Present results for main outcomes, indicating the number of studies and participants for each.", "Yes: Results (direction counts for comprehension, efficiency, writing outcomes)."],
    ["9", "Limitations of evidence", "Provide a brief summary of the limitations of the evidence.", "Yes: Results (risk of bias of learning outcomes; very low certainty)."],
    ["10", "Interpretation", "Provide a general interpretation of the results and important implications.", "Yes: Conclusions."],
    ["11", "Funding", "Specify the primary source of funding for the review.", "Reported in the Funding statement; not included in the abstract under the journal format."],
    ["12", "Registration", "Provide the register name and registration number.", "Yes: Methods (PROSPERO CRD420261513927)."],
]


def checklist():
    doc = base_doc(landscape=True)
    p = doc.add_paragraph()
    add_runs(p, [("PRISMA 2020 Checklist", {"b": True})], size=14)
    p = doc.add_paragraph()
    add_runs(p, [(TITLE, {"i": True})], size=10.5)
    p = doc.add_paragraph()
    add_runs(p, [("Submission to Healthcare (MDPI). Item wording is a concise paraphrase of the PRISMA 2020 checklist "
                  "(Page et al., BMJ 2021;372:n71; CC BY 4.0). Locations refer to section, table and figure numbers of the "
                  "submitted manuscript and Supplementary Materials, so they remain valid after typesetting.", {})], size=9.5)
    p = doc.add_paragraph()
    add_runs(p, [("Table 1. ", {"b": True}), ("PRISMA 2020 main checklist.", {})], size=10)
    grp = {i for i, r in enumerate(PRISMA) if len(r) == 1}
    simple_table(doc, ["Section and topic", "Item", "Checklist item", "Location where item is reported"], PRISMA,
                 [4.2, 1.2, 11.0, 9.2], group_rows=grp)
    doc.add_paragraph()
    p = doc.add_paragraph()
    add_runs(p, [("Table 2. ", {"b": True}), ("PRISMA 2020 for Abstracts checklist.", {})], size=10)
    simple_table(doc, ["Item", "Topic", "Checklist item", "Reported in the abstract?"], PRISMA_A, [1.2, 3.6, 10.8, 10.0])
    p = doc.add_paragraph()
    add_runs(p, [("Reference: Page, M.J.; McKenzie, J.E.; Bossuyt, P.M.; et al. The PRISMA 2020 statement: An updated "
                  "guideline for reporting systematic reviews. BMJ 2021, 372, n71. https://doi.org/10.1136/bmj.n71.", {})], size=8.5)
    doc.core_properties.title = "PRISMA 2020 Checklist"
    normalise(doc)
    doc.save(os.path.join(OUT, "04_PRISMA_2020_Checklist.docx"))


COVER = [
    ("date", "26 September 2026"),
    ("p", "Editor-in-Chief and Editorial Board"),
    ("p", "*Healthcare* (MDPI)"),
    ("p", "Dear Editor,"),
    ("p", f"We submit our systematic review, “{TITLE},” for consideration as a Systematic Review in *Healthcare*."),
    ("h", "Why this review matters"),
    ("p", "Almost one child in ten, about 240 million worldwide, lives with a disability, and World Health Assembly "
          "resolution WHA71.8 commits health systems to improving access to assistive technology as part of universal "
          "health coverage. Speech-to-text, text-to-speech, automatic captioning, and AI-assisted Braille tools are now "
          "prescribed and procured at scale for children with dyslexia, intellectual disability, hearing loss, and visual "
          "impairment, but these decisions often rest on technical accuracy rather than on evidence about the child’s "
          "functioning. Rehabilitation professionals, school health teams, and assistive-technology services need to know "
          "what functional benefit to expect, for whom, and with what confidence."),
    ("h", "What the review adds"),
    ("b", "To our knowledge, it is the first review to evaluate all five speech and Braille functions within one outcome framework mapped to "
          "the WHO International Classification of Functioning, Disability and Health for Children and Youth (ICF-CY), "
          "and it presents the result as an evidence and gap map."),
    ("b", "It separates **access** effects (performance while the technology is in use) from **learning** effects (skills "
          "that persist without it). This distinction reveals a clear and practically important pattern, reflected in the title. "
          "Text-to-speech reduced reading time or effort in five of six reports, and speech-to-text increased text length "
          "and reduced residual errors, "
          "whereas every favourable effect on durable skills or attainment came from studies at high or serious risk of bias."),
    ("b", "It applies design-specific appraisal to randomised, non-randomised, and single-case studies (RoB 2, ROBINS-I, and "
          "What Works Clearinghouse standards), grades certainty with GRADE, and displays all 27 reports in an effect-direction plot."),
    ("b", "It translates the findings into a five-step monitored-trial approach for practitioners and a prioritised "
          "research agenda. It also identifies two functions, automatic classroom captioning and optical Braille recognition, "
          "that have no learner-level evaluation at all, and notes that 24 of 27 reports came from high-income countries."),
    ("h", "Fit with Healthcare"),
    ("p", "The review addresses assistive-technology provision, paediatric rehabilitation, and digital health technology, "
          "and it has direct implications for service delivery, commissioning, and health equity. We believe it will interest "
          "readers working in rehabilitation, school health, digital health, and health policy."),
    ("h", "Rigour and transparency"),
    ("p", "The review was prospectively registered (PROSPERO CRD420261513927) and is reported according to PRISMA 2020, "
          "PRISMA-S, and SWiM; the completed PRISMA 2020 checklist, including the checklist for abstracts, is provided. All "
          "extraction, appraisal, certainty, and effect-direction coding tables are included as Supplementary Materials. "
          "Limitations are stated plainly in the manuscript, including single-reviewer screening, extraction, and appraisal, "
          "and the search sources that were not covered. A post-search check before submission identified five potentially "
          "eligible reports that were not included; they are listed as studies awaiting classification, and all five are "
          "consistent with the review’s conclusions."),
    ("h", "Declarations"),
    ("p", "The manuscript is original, has not been published, and is not under consideration elsewhere. An earlier version "
          "was submitted to *Children* (MDPI) and was declined without external review; it has since been substantially "
          "restructured and reframed for a health and rehabilitation readership. Both authors approved the submission. "
          "Ethics approval was not required because the review used published aggregate data. The work was funded by the "
          "King Salman Center for Disability Research (KSCDR-ERC-2024-01). The authors declare no conflicts of interest. The "
          "use of a generative AI tool in preparing the manuscript is disclosed in the Acknowledgments."),
    ("p", "Thank you for considering our manuscript."),
    ("p", "Sincerely,"),
    ("sig", "Ashit Kumar Dutta (corresponding author), on behalf of Ashit Kumar Dutta and Nasser Ali Aljarallah"),
    ("p", "Department of Computer Science and Information Systems, College of Medical Sciences, AlMaarefa University, "
          "Riyadh, Saudi Arabia; King Salman Center for Disability Research, Riyadh, Saudi Arabia"),
    ("p", "Email: adotta@um.edu.sa"),
]


def cover():
    doc = base_doc()
    inline = Inline({}, [])
    for kind, text in COVER:
        if kind == "h":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            add_runs(p, [(text, {"b": True})], size=11)
        elif kind == "b":
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, inline.parse(text), size=10.5)
        elif kind == "date":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_runs(p, [(text, {})], size=10.5)
        elif kind == "sig":
            p = doc.add_paragraph()
            add_runs(p, [(text, {"b": True})], size=10.5)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, inline.parse(text), size=10.5)
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.name = FONT
    doc.core_properties.title = "Cover letter"
    normalise(doc)
    doc.save(os.path.join(OUT, "01_Cover_Letter_Healthcare.docx"))


if __name__ == "__main__":
    checklist()
    cover()
    print("saved checklist and cover letter")
