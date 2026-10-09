"""Small helpers shared by the manuscript and log builders (python-docx, original template styles)."""
import re
import copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TOKEN = re.compile(r"(\[\[.*?\]\]|\*\*.*?\*\*|__.*?__)", re.S)


def add_runs(par, text, size=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("[[") and part.endswith("]]"):
            r = par.add_run("[" + part[2:-2] + "]")
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("__") and part.endswith("__"):
            r = par.add_run(part[2:-2])
            r.italic = True
        else:
            r = par.add_run(part)
        if size:
            r.font.size = Pt(size)
    return par


def para(doc, text, style="Normal", size=None):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, size)
    return p


def bullet(doc, text, size=None):
    p = doc.add_paragraph(style="List Bullet")
    add_runs(p, text, size)
    return p


def _cell_props(cell, width, header):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("w:tcW", "w:shd", "w:tcBorders", "w:tcMar", "w:vAlign"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    tcW = OxmlElement("w:tcW"); tcW.set(qn("w:type"), "dxa"); tcW.set(qn("w:w"), str(width)); tcPr.append(tcW)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement("w:" + side); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:color"), "BFBFBF")
        borders.append(b)
    tcPr.append(borders)
    if header:
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "E8E8E8")
        tcPr.append(shd)
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        m = OxmlElement("w:" + side); m.set(qn("w:w"), "70"); m.set(qn("w:type"), "dxa"); mar.append(m)
    tcPr.append(mar)


def table(doc, headers, rows, widths, size=9.5):
    """widths in DXA; must sum to the usable width (9360 for 6.5 in)."""
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    look = tblPr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(layout)
    else:
        tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa"); tblW.set(qn("w:w"), str(sum(widths)))
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); trPr.append(th)
    for ri, row in enumerate([headers] + rows):
        for ci, txt in enumerate(row):
            cell = t.rows[ri].cells[ci]
            _cell_props(cell, widths[ci], ri == 0)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            lines = txt.split("\n")
            for li, line in enumerate(lines):
                if li:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.0
                add_runs(p, ("**" + line + "**") if ri == 0 and line else line, size)
    return t


def clear_body_after(doc, keep_first=1):
    body = doc.element.body
    children = [c for c in body.iterchildren() if c.tag != qn("w:sectPr")]
    for c in children[keep_first:]:
        body.remove(c)
