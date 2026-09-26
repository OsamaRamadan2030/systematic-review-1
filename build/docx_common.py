"""Shared helpers for building MDPI-styled Word documents with python-docx."""
import copy
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

TOKEN_RE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|\^[^^]+?\^|~[^~\s]+?~|\[@[^\]]+\])")


def compress(nums):
    """[1,2,3,5,7,8] -> '1–3,5,7,8' (MDPI style: ranges for 3+ consecutive)."""
    nums = sorted(set(nums))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            out.append(f"{nums[i]}–{nums[j]}")
        else:
            out.extend(str(n) for n in nums[i:j + 1])
        i = j + 1
    return ",".join(out)


class Inline:
    """Parses inline markup into (text, fmt) segments; resolves [@key] citations."""

    def __init__(self, numbers, included_order):
        self.numbers = numbers
        self.included = included_order

    def cite(self, body):
        keys = []
        for k in body.split(";"):
            k = k.strip().lstrip("@")
            keys.extend(self.included if k == "INCLUDED" else [k])
        return "[" + compress(self.numbers[k] for k in keys) + "]"

    def parse(self, text, fmt=None):
        fmt = dict(fmt or {})
        segs = []
        pos = 0
        for m in TOKEN_RE.finditer(text):
            if m.start() > pos:
                segs.append((text[pos:m.start()], fmt))
            tok = m.group(0)
            if tok.startswith("**"):
                segs.extend(self.parse(tok[2:-2], {**fmt, "b": True}))
            elif tok.startswith("[@"):
                segs.append((self.cite(tok[2:-1]), fmt))
            elif tok.startswith("*"):
                segs.extend(self.parse(tok[1:-1], {**fmt, "i": True}))
            elif tok.startswith("^"):
                segs.append((tok[1:-1], {**fmt, "sup": True}))
            elif tok.startswith("~"):
                segs.append((tok[1:-1], {**fmt, "sub": True}))
            pos = m.end()
        if pos < len(text):
            segs.append((text[pos:], fmt))
        return segs

    def plain(self, text):
        return "".join(t for t, _ in self.parse(text))


def add_runs(par, segs, size=None, color=None):
    for text, fmt in segs:
        r = par.add_run(text)
        if fmt.get("b"):
            r.bold = True
        if fmt.get("i"):
            r.italic = True
        if fmt.get("sup"):
            r.font.superscript = True
        if fmt.get("sub"):
            r.font.subscript = True
        if size:
            r.font.size = Pt(size)
    return par


def set_cell_borders(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, sz in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)


def set_cell_margins(cell, top=30, bottom=30, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def three_line_table(doc, header, rows, widths, inline, style="MDPI_4.2_table_body", size=8,
                     group_rows=None, align_left_cols=None, indent_twips=None, raw=False):
    """MDPI three-line table: top and bottom rules, header rule, fixed widths (twips).

    rows: list of lists of markup strings; group_rows: set of row indices rendered as
    full-width bold subgroup headers (row given as [text]).
    """
    group_rows = group_rows or set()
    align_left_cols = set(range(len(widths))) if align_left_cols is None else set(align_left_cols)
    table = doc.add_table(rows=1 + len(rows), cols=len(widths))
    tbl = table._tbl
    tblPr = tbl.tblPr
    for tag in ("w:tblStyle",):
        el = tblPr.find(qn(tag))
        if el is not None:
            tblPr.remove(el)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    if indent_twips is not None:
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), str(indent_twips))
        ind.set(qn("w:type"), "dxa")
        tblPr.append(ind)
    else:
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        tblPr.append(jc)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))

    def fill(cell, text, width, bold=False, left=True):
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW")
            tcPr.insert(0, tcW)
        tcW.set(qn("w:w"), str(width))
        tcW.set(qn("w:type"), "dxa")
        set_cell_margins(cell)
        lines = text.split("\n") if text else [""]
        p = cell.paragraphs[0]
        for li, line in enumerate(lines):
            if li > 0:
                p = cell.add_paragraph()
            p.style = doc.styles[style]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
            segs = [(line, {"b": bold})] if raw else inline.parse(line, {"b": bold} if bold else None)
            add_runs(p, segs, size=size)

    hdr = table.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trPr.append(th)
    for j, (cell, text) in enumerate(zip(hdr.cells, header)):
        fill(cell, text, widths[j], bold=True, left=j in align_left_cols)
        set_cell_borders(cell, bottom=4)
    for i, row in enumerate(rows):
        tr = table.rows[i + 1]
        cant = OxmlElement("w:cantSplit")
        tr._tr.get_or_add_trPr().append(cant)
        if i in group_rows:
            merged = tr.cells[0].merge(tr.cells[-1])
            fill(merged, row[0], sum(widths), bold=True, left=True)
            set_cell_shading(merged, "F2F2F2")
            for para in merged.paragraphs:  # keep a subgroup header on the same page as its first row
                para.paragraph_format.keep_with_next = True
            continue
        for j, (cell, text) in enumerate(zip(tr.cells, row)):
            fill(cell, text, widths[j], left=j in align_left_cols)
    return table


def section_break_paragraph(doc, sectPr_template):
    """Append an empty paragraph that ends the current section with a copy of sectPr_template."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.append(copy.deepcopy(sectPr_template))
    return p


def remove_journal_logo(doc, marker="Children journal logo"):
    """Remove header drawings whose alt text names the previous journal's logo. Returns count removed."""
    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    removed = 0
    for rel in list(doc.part.rels.values()):
        if "header" not in rel.reltype:
            continue
        el = rel.target_part.element
        for drawing in list(el.iter(qn("w:drawing"))):
            if any(marker in (dp.get("descr") or "") for dp in drawing.iter(wp + "docPr")):
                run = drawing.getparent()
                run.getparent().remove(run)
                removed += 1
    return removed


_TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
                "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
                "tblCaption", "tblDescription"]
_TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"]
_TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight",
               "tblHeader", "tblCellSpacing", "jc", "hidden"]


def _reorder(el, order):
    local = lambda e: e.tag.split("}")[1]
    kids = list(el)
    known = [k for k in kids if local(k) in order]
    other = [k for k in kids if local(k) not in order]
    # drop duplicates of singleton elements, keeping the last occurrence
    seen, dedup = {}, []
    for k in known:
        seen[local(k)] = k
    dedup = sorted(seen.values(), key=lambda e: order.index(local(e)))
    for k in kids:
        el.remove(k)
    for k in dedup + other:
        el.append(k)


def normalise(doc):
    """Reorder table/row/cell property children into schema order and fix settings zoom."""
    body = doc.element.body
    for tag, order in (("w:tblPr", _TBLPR_ORDER), ("w:tcPr", _TCPR_ORDER), ("w:trPr", _TRPR_ORDER)):
        for el in body.iter(qn(tag)):
            _reorder(el, order)
    settings = doc.settings.element
    for z in settings.iter(qn("w:zoom")):
        if z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")
