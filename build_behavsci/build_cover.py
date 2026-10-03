import docx, re
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
TITLE = re.search(r'^@TITLE (.*)$', open('/home/user/build/manuscript.txt').read(), re.M).group(1)
d = docx.Document(); st = d.styles['Normal']; st.font.name = 'Palatino Linotype'; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Palatino Linotype')
s = d.sections[0]; s.page_width, s.page_height = Cm(21.0), Cm(29.7)
for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(s, m, Cm(2.5))
def P(t, b=False, after=8, italic_title=False):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    for part in re.split(r'(\*\*[^*]+\*\*)', t):
        if part: r = p.add_run(part.strip('*') if part.startswith('**') else part); r.bold = b or part.startswith('**')
    return p
P('3 October 2026')
P('Dr. Megan Israelsen-Augenstein\nGuest Editor, Special Issue “Advancing Evidence-Based Evaluation and Treatment of Children with Language and Reading Disorders”\n*Behavioral Sciences*'.replace('*',''), after=12)
P('Dear Dr. Israelsen-Augenstein,')
P(f'We submit our manuscript, “{TITLE}”, for consideration as a Systematic Review in the Special Issue “Advancing Evidence-Based Evaluation and Treatment of Children with Language and Reading Disorders”.')
P('Text-to-speech, speech-to-text, and related technologies are routinely provided to children with reading and writing disorders, yet studies often report performance while a tool is in use as though it were evidence of learning. Our review addresses this evaluation problem directly. Across 47 reports, we classify every outcome by its measurement condition, distinguishing assisted performance, unaided learning, and proximal acquisition of trained content, and we synthesize findings within these categories.')
P('The main findings are clinically and educationally relevant. Speech-to-text often increased the amount written, but accuracy and quality varied. Text-to-speech comprehension effects ranged from favorable to null or favoring the comparator. Only 6 of the 47 reports tested unaided skills in a controlled comparison, and their results were inconsistent; one randomized experiment found poorer short-term orthographic learning after reading with text-to-speech than after independent reading. The review translates these findings into a practical recommendation: set and monitor access goals and learning goals separately, using brief within-child comparisons.')
P('The work fits the Special Issue’s focus on evidence-based evaluation and technology-assisted support for children with language and reading difficulties. Most included learners had dyslexia or other reading and writing disabilities. Evidence on sensory, motor, and intellectual disabilities is retained to examine access across literacy modalities, without relabeling those conditions as language disorders.')
P('The review was registered in PROSPERO (CRD420261513927) and is reported according to PRISMA 2020, PRISMA-S, and SWiM; the completed checklists and a PRISMA flow diagram are included. We have reported its limitations openly, including single-reviewer screening and extraction, post hoc refinement of the outcome framework, reports verified only through abstracts, and the use of a generative AI tool, which is disclosed in the Methods and Acknowledgments.')
P('Earlier versions of this work were submitted to *Healthcare* and *Children* and were declined. The manuscript has since been substantially revised, including a supplementary ERIC search, verification of all included reports against their original sources, correction of eligibility and extraction errors, and a new synthesis framework.'.replace('*',''))
P('This manuscript is original, has not been published, and is not under consideration elsewhere. Both authors have approved the submission. The review was funded by the King Salman Center for Disability Research, with which both authors are affiliated; the funder had no role in the review. There are no other competing interests.')
P('Thank you for considering our manuscript.', after=14)
P('Sincerely,', after=2)
P('Ashit Kumar Dutta, on behalf of both authors\nDepartment of Computer Science and Information Systems, College of Applied Sciences, AlMaarefa University, Riyadh, Saudi Arabia; King Salman Center for Disability Research, Riyadh, Saudi Arabia\nadotta@um.edu.sa')
d.save('/home/user/build/out/Cover_Letter.docx'); print('ok')
