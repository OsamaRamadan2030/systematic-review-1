import re
import sys
from docx import Document
from docxlib import para, table, bullet, clear_body_after

SRC, MANU, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
doc = Document(SRC)
clear_body_after(doc, keep_first=1)
doc.paragraphs[0].runs[0].text = "Author Verification and Revision Log: Revised Introduction and Methods"
for r in doc.paragraphs[0].runs[1:]:
    r.text = ""

H1 = lambda t: para(doc, t, "Heading 1")
H2 = lambda t: para(doc, t, "Heading 2")
P = lambda t: para(doc, t)
B = lambda t: bullet(doc, t)

P("Companion to Revised_Introduction_and_Methods.docx. This log is for the authors and is not for submission. It records "
  "what was changed, how each item in Essential Author Checks was handled, every value that only the study records can "
  "supply, and the checks to complete before submission.")

# ------------------------------------------------------------------ 1
H1("1 How to use the revised manuscript")
B("Text highlighted in yellow and enclosed in square brackets is a field that only the original study records can "
  "complete. No number, date, criterion or procedure in those fields has been invented. Replace each field with the "
  "verified fact, or delete the sentence if the procedure did not occur, and remove the highlight.")
B("Unhighlighted text restates what the supplied draft and records reported, applies standard methodology to the "
  "procedures already described or gives arithmetic derived from stated values (Section 5). Read every sentence against "
  "the study records before submission, because a Methods section must describe what was done, not what should have "
  "been done.")
B("Nothing highlighted may remain at submission. Section 7 lists every field, taken directly from the manuscript.")

# ------------------------------------------------------------------ 2
H1("2 Scope decisions")
B("**Title, design and Introduction.** The title, the three-arm nonrandomised quasi-experimental design, the completed-study "
  "wording and the Introduction are unchanged.")
B("**Instruments and tools were not replaced.** The study is complete. Replacing the simulator, rubric, knowledge test, "
  "checklist, AI model or any questionnaire in the report would describe a study that did not take place. That is "
  "misreporting, and reviewers usually detect it from inconsistencies between Methods, Results and the records. The "
  "Methods therefore keep every tool actually used. They strengthen the account of how each tool was developed, scored, "
  "checked and analysed, and they state its limits openly. This is how a weakness becomes defensible.")
B("**Omitted from this report, as authorised in Essential Author Checks:** the cognitive-load questionnaire paragraph, "
  "results and licence statement, and the pilot nurse–student known-groups comparison. Neither is replaced. Section 2.1 "
  "acknowledges that additional exploratory measures were collected.")
B("**Length.** The Methods section is about 4,700 words, plus four tables. If the target journal's limit is tighter, move "
  "Section 2.10 and Tables 3 and 4 to supplementary files. Cite a supplementary file only after it has been prepared.")

# ------------------------------------------------------------------ 3
H1("3 Essential Author Checks: resolution")
W = [2000, 3900, 3460]
table(doc, ["Check", "How the revised Methods handle it", "What the authors must do"], [
    ["1a Simulator and infant cases",
     "Brand and model names removed. Section 2.9.1 now requires an explicit statement of the simulator and represented age for each age band and of how fontanelle and other infant findings were presented and credited. A generic description can no longer hide the mismatch.",
     "Confirm the equipment from the laboratory inventory. Do not substitute a different model. Use Option A or B in Section 6.1, and if Option B applies, add the limitation to the Discussion."],
    ["1b PN-CRER scoring rule",
     "Section 2.9.2 removes the conflicting 'observation-only' wording. It now requires the evidence source for each domain and the caps for stated-only actions and critical errors, exactly as in the final manual.",
     "Supply the final scoring manual and case keys. Copy the rule applied to the recordings and confirm that every recording was scored under it."],
    ["2 Knowledge-test equating",
     "Section 2.9.3 specifies the method that creates a single metric: random-groups linear equating of Form B onto reference Form A, with Form A untransformed and test-form sequence included in the sensitivity analyses. Raw 0–40 bounds are retained.",
     "The supplied transformations are near inverses of each other (Section 5.2). If both forms were transformed in the completed analysis, the scores are not on one metric. Re-equate, re-run the knowledge analyses and disclose the correction (Section 6.2). Do not silently substitute raw scores."],
    ["3a Content validity indices",
     "Section 2.10.1 defines the I-CVI, S-CVI/Ave, S-CVI/UA and modified kappa with acceptability criteria, computed from the original rating forms. No values appear in the Methods.",
     "Authenticate the expert worksheets. If every rating was 3 or 4, all indices equal 1.00, and the reported 0.86, 0.82 and 0.94 cannot stand. Report recalculated values in the Results."],
    ["3b Generalisability study",
     "Section 2.10.3 specifies a REML crossed random-effects model for the incompletely crossed design, with the explicit relative and dependability formulas for the operational design.",
     "Obtain the person–case matrix, the fitted model and the coefficient code. The supplied components imply about 0.673, not 0.82, for two cases and one rater (Section 5.4). Report only recomputed values."],
    ["3c Nurse–student comparison",
     "Excluded from the Methods, as authorised.",
     "If reinstated, the supplied summaries imply pooled SD 3.982, d 1.090 and t(30) 2.669, p = .012, not the supplied values. Recompute from the raw data."],
    ["3d Agreement models and LCJR count",
     "Changing rater pairs: one-way ICC. Common four-rater subset: two-way ICC. Cluster bootstrap by student. LCJR stated as 270 primary-rated and 135 double-rated recordings. Double-rating counts given without percentages.",
     "Confirm the model actually fitted. Add Results values only after reconciliation, each with its n and 95% CI."],
    ["3e Repeated recordings in correlations",
     "Section 2.10.3 makes the student the sampling unit for the PN-CRER–LCJR association, with cluster-bootstrap CIs.",
     "Choose and confirm the method (repeated-measures correlation or multilevel model)."],
    ["4a Protocol history",
     "Section 2.1 and Table 4 require the order of plan finalisation, rating completion, database lock and unmasking, and a list of changes between versions. The study is described as not registered, and only version-1 analyses are called prespecified.",
     "Supply time-stamped records and use Option A or B in Section 6.3."],
    ["4b Reporting scope and omitted measures",
     "Section 2.1 acknowledges the additional measures and requires a statement that the decision to report them separately did not depend on the results.",
     "Name the measures. If any was a protocol outcome, disclose its omission and the reason. Check that no other document describes them as never collected."],
    ["4c Cognitive-load questionnaire",
     "Omitted from this report, not replaced and not described as never collected.",
     "Keep it in the protocol-to-report record. Make no licence claim."],
    ["5 Cross-references",
     "Tables 1–4 and Figure 1 are each cited in the text. No supplementary file is cited. Figure counts are 144 + 142 + 143 = 429 of 450.",
     "Match all counts to the Results and analysis dataset. Complete the per-arm attendance and missingness fields in Figure 1."],
    ["6 References",
     "Original references and published titles retained. Titles mentioning randomised trials describe other studies. Ten methodological references added after their identity was checked (Section 4).",
     "Check every DOI against the journal style before submission."],
    ["7 Theory and figures",
     "The framework remains a short paragraph (Section 2.5). No theoretical figure was added. The participant-flow figure was retained and expanded.",
     "None."],
    ["8 Instrument ownership and permissions",
     "PN-CRER, PNNKT and PNPSC are described as investigator-developed. Section 2.13 now requires their owner and sharing terms. LCJR is described as used with the author's permission.",
     "Complete the owner and terms fields and keep the LCJR permission letter on file."],
    ["9 Generative AI declaration",
     "A declaration is drafted after the Figure, separate from the GPT-4o intervention.",
     "Name only the tools actually used in drafting (ChatGPT for earlier drafts if applicable, and Claude for this revision). Finalise it only after the authors have reviewed the manuscript."],
], W)

# ------------------------------------------------------------------ 4
H1("4 Additional weaknesses identified and corrected")
W3 = [2700, 4300, 2360]
table(doc, ["Weakness in the previous Methods", "Correction", "Section"], [
    ["No defined estimand; handling of non-attendance and contamination implicit", "Primary estimand defined, with a treatment-policy strategy for attendance, cross-arm material use and outside AI use", "2.11.1"],
    ["Planning assumptions with no implied precision", "Design effect, power and minimum detectable difference derived from the stated assumptions and labelled as derived", "2.3"],
    ["Allocation variables not linked to the analysis", "Facilitator identified as the allocation stratum. All balancing variables are adjusted for, and timetable position is added in sensitivity analysis", "2.4, 2.11.2, 2.11.5"],
    ["Group formation and facilitator or instructor characteristics unstated", "Fields added for the group-formation rule, qualifications and rotation of simulation instructors across arms", "2.2, 2.7"],
    ["Session segments not identified as common or arm-specific", "Arm-specific segments specified; consolidation and product rows added to Table 1", "2.7, Table 1"],
    ["How much of the intervention was actually AI content was unclear", "Proportions of unchanged AI, modified AI and expert-added nodes reported; candidate-rating criteria and prompt retention specified", "2.6"],
    ["Comparators' feedback materials not quality-matched", "Field to confirm that the clinical panel also reviewed facilitator guides for all arms", "2.6"],
    ["Masking of assessment staff and statistician not stated", "Fields added", "2.8"],
    ["Fidelity sampling and coder agreement unspecified", "24 of 75 sessions, sampling method, checklist content and kappa specified", "2.8"],
    ["Contamination questionnaire could be read as a scale", "Items analysed singly and not summed", "2.8"],
    ["Test security after piloting with students from the same faculty", "Confidentiality field added; advance case knowledge collected by student report", "2.10.2, 2.8"],
    ["Item retention criteria, rater calibration criterion, rating period and drift frequency missing", "Fields added", "2.10.2, 2.10.3"],
    ["No standardised effect size or ICC reporting", "Hedges-type standardisation from an unconditional model; section and group ICCs reported", "2.11.2"],
    ["Baseline comparison method unstated", "Standardised mean differences, without significance tests", "2.11.6"],
    ["Imputation method, number of imputations and handling of withdrawals unstated", "Arm-wise multilevel imputation, Rubin's rules, delta adjustment and fields for the remaining details", "2.11.4"],
    ["Software packages and response to diagnostic problems unstated", "Fields added", "2.11.6"],
    ["Database lock, data entry checking and retention period unstated", "Fields added", "2.13"],
    ["GREET not cited for an educational intervention", "Added", "2.1"],
    ["No procedural timeline", "Table 4 added", "Table 4"],
    ["Flow diagram lacked exposure and per-arm reasons", "Figure 1 redrawn with attendance, reasons and both analysis populations", "Figure 1"],
    ["Validity evidence scattered", "Table 3 maps evidence to the scoring, generalisation and extrapolation inferences", "2.10, Table 3"],
], W3)
P("Added references, each checked against a bibliographic database: Bates et al. (2015); Briesch et al. (2014); Cook et "
  "al. (2015); Cro et al. (2020); Hedges (2007); Hothorn et al. (2008); International Council for Harmonisation (2019); "
  "Lynn (1986); Phillips et al. (2016); Watts et al. (2021). Confirm the full author list for Watts et al. (2021) and "
  "remove any reference whose method was not used, for example Watts et al. if the simulation design standard was not "
  "followed.")

# ------------------------------------------------------------------ 5
H1("5 Verification calculations")
H2("5.1 Planning precision (Section 2.3)")
P("Design effect for 27 analysed students per section in groups averaging 4.5: 1 + (4.5 − 1)(0.03 + 0.05) + "
  "(27 − 4.5)(0.03) = 1.955. Baseline adjustment with r = 0.50 multiplies residual variance by 0.75. Effective sample "
  "size per arm: 135 / 1.955 / 0.75 = 92.1. Standard error of the standardised difference: √(2 / 92.1) = 0.147. Power for "
  "δ = 0.50 with a noncentral t distribution on 8 degrees of freedom: 0.84 (0.92 with a normal reference). Minimum "
  "detectable difference at 80% power: 0.47 SD, or 2.4 points. With no loss (150 per arm, design effect 2.07), the power is "
  "0.86. These values were calculated for this revision, not taken from the protocol.")
H2("5.2 Equating transformations")
P("The inverse of 0.97 × A + 1.24 is (x − 1.24) / 0.97 = 1.031x − 1.278, which matches the supplied 1.03 × B − 1.28. "
  "Each transformation therefore maps one form onto the other form's scale. Applying the first to Form A and the second to "
  "Form B swaps their metrics instead of placing both on one reference metric. Calculated endpoints are 12.88–38.10 (A, raw "
  "12–38) and 13.14–38.89 (B, raw 14–39), which do not match the recorded 11.8–37.9 and 13.9–38.7. The analysed knowledge "
  "variable must therefore be identified from the analysis code.")
H2("5.3 Content validity")
P("With eight experts and every rating 3 or 4, each I-CVI is 8/8 = 1.00. S-CVI/Ave and S-CVI/UA are then 1.00, and the "
  "modified kappa is 1.00 because the probability of chance agreement is 0.5⁸ ≈ 0.004. Values of 0.86, 0.82 or 0.94 "
  "require ratings of 1 or 2 that are absent from the supplied matrix.")
H2("5.4 Generalisability coefficient")
P("Relative coefficient for n′c cases and n′r raters: σ²(p) / [σ²(p) + σ²(pc)/n′c + σ²(pr)/n′r + σ²(pcr,e)/(n′c n′r)]. "
  "Essential Author Checks reports that the supplied components give about 0.673 for n′c = 2 and n′r = 1, against a recorded "
  "0.82. Recompute from the fitted model, and state whether a student's two cases were scored by the same rater, because "
  "that determines n′r.")
H2("5.5 Pilot nurse–student comparison (excluded)")
P("t = 2.669 on 30 degrees of freedom gives a two-sided p of 0.0122, consistent with the recalculation in Essential "
  "Author Checks.")
H2("5.6 Dates and contact time")
P("Session minutes: 15 + 30 + 30 + 15 + 10 + 50 + 20 + 10 = 180. Contact time: 30 + 4 × 180 = 750 minutes. Preparation: "
  "4 × 60 = 240 minutes. Allocations: 3! for each of five facilitators gives 6⁵ = 7,776. Assessment interval: if final "
  "sessions fell on 15–17 March, 29 March–7 April gives 12–23 days, as stated, and the 8 April knowledge test was 22–24 "
  "days after the final session. Fidelity sample: 5 sessions × 5 sections = 25 per arm, of which 8 per arm (24 of 75) were "
  "coded.")
H2("5.7 Calendar context")
P("The teaching period, 22 February–17 March 2026, overlapped Ramadan (approximately 18 February–19 March 2026), and Eid "
  "al-Fitr (about 20 March) fell before follow-up. This affected all arms equally, but it may explain the holiday-travel "
  "losses and could have affected attendance or engagement. Consider stating it in the Discussion, and add it to the Methods "
  "only if the session schedule was adjusted.")

# ------------------------------------------------------------------ 6
H1("6 Alternative wordings for unresolved records")
H2("6.1 Simulators (Section 2.9.1)")
P("**Option A: an age-appropriate simulator was used for the young-child cases.** Young-child cases used a high-fidelity "
  "infant simulator representing a child aged [age], and school-age cases used a high-fidelity pediatric simulator "
  "representing a five-year-old child. Findings that the infant simulator could not physically reproduce, including "
  "[findings], were presented by [method] at fixed times, and case keys credited the corresponding assessment.")
P("**Option B: the five-year-old simulator was used for all cases.** All cases used a high-fidelity pediatric simulator "
  "representing a five-year-old child, because [reason]. For young-child cases, the scenario stated the child's age and "
  "weight, age-appropriate equipment sizes were supplied and monitor parameters were set to age-appropriate values. "
  "Findings that the simulator could not represent, including anterior fontanelle findings, were presented by [scripted "
  "standardised-parent report, cue card or monitor text] at fixed times. Case keys credited students for performing or "
  "requesting the corresponding assessment. Add to the Discussion: the physical representation of the young-child cases "
  "did not match the represented age, which limits assessment fidelity for infant-specific examination findings.")
H2("6.2 Equating (Section 2.9.3)")
P("**If the records confirm a single reference metric:** complete the highlighted equating fields in the manuscript.")
P("**If both forms were transformed:** re-equate Form B to Form A from the baseline random groups, repeat all knowledge "
  "analyses and add to Section 2.9.3: \"During manuscript preparation, an error was identified in the original equating, "
  "in which each form had been transformed towards the other form's scale. Knowledge scores were re-equated to the Form A "
  "metric, and all knowledge analyses were repeated before submission.\"")
H2("6.3 Analysis-plan sequence (Section 2.1)")
P("**Option A, plan finalised before unmasking:** Version 2 was finalised and time-stamped at [time] on 30 June 2026, "
  "after rating was completed on [date] and the database was locked on [date], and before the strategy code key was "
  "released at [time] on the same day.")
P("**Option B, plan finalised at or after unmasking:** Version 2 was finalised on 30 June 2026, [when] the strategy code key "
  "was released. Changes introduced in version 2 are therefore treated as post hoc; results follow version 1, and analyses "
  "added in version 2 are labelled post hoc.")

# ------------------------------------------------------------------ 7
H1("7 Fields to complete")
P("Every highlighted field in the revised manuscript, in order of appearance.")
m = Document(MANU)
rows, sec = [], ""
for p in m.paragraphs:
    if p.style.name in ("Heading 2", "Heading 3"):
        sec = p.text.split(" ")[0] if p.text[0].isdigit() else p.text
    text = "".join(r.text for r in p.runs)
    hl = []
    cur = ""
    for r in p.runs:
        if r.font.highlight_color is not None:
            cur += r.text
        elif cur:
            hl.append(cur); cur = ""
    if cur:
        hl.append(cur)
    for h in hl:
        rows.append([sec, h.strip("[]").replace("___", "…")])
for t in m.tables:
    title = None
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if r.font.highlight_color is not None:
                        rows.append(["Table: " + t.rows[0].cells[0].text + " / " + row.cells[0].text, r.text.strip("[]").replace("___", "…")])
seen, uniq = set(), []
for r in rows:
    k = tuple(r)
    if k not in seen:
        seen.add(k); uniq.append(r)
rows = uniq
rows.append(["Figure 1", "Per-arm attendance (≥3, 1–2, 0 sessions) and per-arm reasons for missing follow-up"])
table(doc, ["Location", "Field"], rows, [2300, 7060], size=9)
P("Total fields: %d." % len(rows))

# ------------------------------------------------------------------ 8
H1("8 Pre-submission checklist")
for t in [
    "All highlighted fields completed or their sentences deleted; no square brackets remain.",
    "Every number in the Results matches the analysis output and the Methods (450 allocated; 150 per arm; 144, 142 and 143 followed up; 429 in total).",
    "Content validity, agreement and generalisability values recalculated from the original records before they appear in the Results, each with n and 95% CI.",
    "Knowledge analyses confirmed or re-run on a single equated metric.",
    "Simulator wording confirmed from the inventory, with a limitation added if Option B applies.",
    "Analysis-plan sequence documented; post hoc analyses labelled as post hoc.",
    "TREND checklist completed with page numbers; TIDieR and GREET items cross-checked against Table 1 and Sections 2.6–2.8.",
    "Supplementary files cited only if they will be uploaded.",
    "Instrument ownership and sharing terms stated; LCJR permission on file.",
    "Generative AI declaration finalised and naming only the tools actually used.",
    "References formatted to the target journal's style and DOIs checked.",
    "Word count checked against the journal limit; Section 2.10 and Tables 3–4 moved to supplementary files if needed.",
]:
    B(t)

doc.save(OUT)
print("saved", OUT, len(rows), "fields")
