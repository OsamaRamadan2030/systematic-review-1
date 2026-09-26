"""Build the Healthcare (MDPI) manuscript .docx from manuscript_source.txt.

Uses the authors' MDPI Word template (styles, front-matter box, line numbering,
headers/footers) from original_children_submission/Manuscript_Children.docx,
replaces the body, and removes the Children journal logo from the header.

Usage: python3 build/build_manuscript.py
"""
import copy
import csv
import json
import os
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from docx_common import normalise, Inline, add_runs, section_break_paragraph, three_line_table
from numbering import assign
from references import INCLUDED_ORDER, REFS

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TEMPLATE = os.path.join(ROOT, "original_children_submission", "Manuscript_Children.docx")
SRC = os.path.join(HERE, "manuscript_source.txt")
OUTDIR = os.path.join(ROOT, "submission_healthcare")
OUT = os.path.join(OUTDIR, "02_Manuscript_Healthcare.docx")

STYLE = {
    "p": "MDPI_3.1_text", "pni": "MDPI_3.2_text_no_indent", "h1": "MDPI_2.1_heading1",
    "h2": "MDPI_2.2_heading2", "h3": "MDPI_2.3_heading3", "bullet": "MDPI_3.8_bullet",
    "numbered": "MDPI_3.7_itemize", "abstract": "MDPI_1.7_abstract", "abstract_head": "MDPI_1.7_abstract",
    "keywords": "MDPI_1.8_keywords", "figcap": "MDPI_5.1_figure_caption", "tabcap": "MDPI_4.1_table_caption",
    "tabfoot": "MDPI_4.3_table_footer", "back": "MDPI_6.2_back_matter", "hq": "MDPI_3.2_text_no_indent",
}


# ------------------------------------------------------------------ data
def load_csv(name):
    with open(os.path.join(HERE, "data", name), newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


PARTICIPANTS = {
    "23": "39 LD, 9–18 y (STT 19; control 20)",
    "24": "52 LD, 9–18 y (includes the 39 in [@s23])",
    "25": "21 LD (+10 peers); high school",
    "49": "16 reading/writing difficulties (+12 peers); 10–13 y",
    "26": "8 severe reading/writing difficulties; middle school",
    "27": "7 mild ID; 10–13 y",
    "28": "4 mild ID; 10–13 y",
    "29": "20 dyslexia; 6–12 y (10 per group)",
    "41": "49 reading disability; grades 3–9",
    "30": "149 severe reading/writing difficulties; grades 4 and 8 and high school",
    "31": "29 reading difficulties; 8–12 y",
    "36": "20 specific LD; grades 1–5 (10 per group)",
    "40": "38 dyslexia or ADHD (+21 typical readers); junior high",
    "42": "20 dyslexia (+50 typical); school age",
    "48": "94 reading difficulties; grade 8",
    "32": "4 LD; secondary",
    "33": "2 LD; grade 4",
    "34": "4 reading LD; middle school",
    "35": "3 reading difficulties; grades 3–4",
    "37": "1 ADHD; grade 2",
    "39": "8 mild ID; 11–14 y",
    "43": "2 specific LD; high school",
    "44": "3 reading disabilities; junior high",
    "45": "4 reading LD; secondary",
    "46": "4 LD; 16–17 y",
    "47": "5 mild ID",
    "38": "10 visual impairment; 4 y 11 mo–14 y 11 mo (+7 teachers)",
}
DC = {"RCT": "RCT", "NRS": "NRS", "SCD": "SCED"}


def table1_rows(studies, ed):
    by = {r["ref"]: r for r in studies}
    outcomes = {}
    for e in ed:
        for ref in e["ref"].split(";"):
            outcomes.setdefault(ref, []).append(f"{e['domain']} ({'A' if e['level'] == 'access' else 'L'})")
    groups = [("STT", "Speech-to-text (7 reports; 6 studies)"),
              ("TTS", "Text-to-speech and combined speech technologies (19 reports)"),
              ("BRL", "Adaptive Braille learning (1 report)")]
    rows, grp = [], set()
    for fn, title in groups:
        grp.add(len(rows))
        rows.append([title])
        for key in INCLUDED_ORDER:
            old = key[1:]
            r = by[old]
            if r["function"] != fn:
                continue
            study = f"{r['label']}, {r['year']} [@{key}]; {r['country'].replace('/', ' and ')}"
            design = f"{DC[r['design_class']]}: {r['design']}".replace("(extends [23])", "")
            design = design.replace(" (5 conditions)", "; 5 conditions").strip()
            comp = r["comparator"]
            if len(comp) > 1 and comp[1].islower():
                comp = comp[0].lower() + comp[1:]
            tech = f"{r['technology']} vs. {comp}"
            outc = "; ".join(outcomes.get(old, []))
            appr = f"{r['tool']}: {r['rating']}" if r["rating_class"] != "na" else "Not rated (abstract only)"
            rows.append([study, design, PARTICIPANTS[old], tech, outc, appr])
    return rows, grp


TABLE2 = [
    ["STT-supported writing practice vs. general computer instruction; unaided word recognition, spelling, and reading comprehension [@s23; @s24]",
     "1 study (2 reports); 52; NRS",
     "Direction favoured STT on all measures (*p* < 0.0001 to *p* < 0.04); no effect estimate or CI; legacy discrete-speech software.",
     "⊕◯◯◯ Very low", "RoB (2); Ind (1); Imp (1)"],
    ["STT vs. handwriting, keyboarding, or baseline; text length, transcription accuracy, and text quality [@s25; @s49; @s26; @s27; @s28]",
     "5 studies; 56; NRS 2, SCED 3",
     "Text length favoured STT in 2 of 3 studies and residual errors in 1 of 2 (*F*(1,52) = 14.92, partial η^2^ = 0.22). Text quality favoured STT in 2 of 5 and was unchanged in the counterbalanced keyboard comparison (*p* = 0.95). No pooled estimate or CI.",
     "⊕◯◯◯ Very low", "RoB (2); Inc (1); Ind (1); Imp (1)"],
    ["TTS or combined TTS/STT vs. unsupported reading, usual support, or an alternative accommodation; reading and listening comprehension [@s30; @s31; @s32; @s33; @s34; @s35; @s39; @s40; @s41; @s42; @s43; @s44; @s45; @s46; @s47; @s48]",
     "16 studies; 418; RCT 1, NRS 5, SCED 10",
     "With TTS in use, comprehension favoured TTS in 7 of 15 reports, was mixed or unchanged in 7, and favoured a human reader in 1; benefit varied by grade, reading profile, and comparator. The largest controlled study (*n* = 149) found no between-group difference after 24 sessions or at 1 year. Reading rate, time, or effort favoured TTS in 5 of 6 reports (not separately graded).",
     "⊕◯◯◯ Very low", "Started high (randomised evidence). RoB (1); Inc (2); Ind (1); Imp (1)"],
    ["TTS-supported vs. conventional teaching; general academic achievement [@s29]",
     "1 study; 20; RCT",
     "Post-test mean 28 vs. 24 points on an unvalidated 50-item test (MD 4.0; CI not reported).",
     "⊕◯◯◯ Very low", "Started high. RoB (2); Ind (1); Imp (2)"],
    ["TTS-assisted vs. conventional mathematics instruction; mathematics achievement [@s36]",
     "1 study; 20; NRS",
     "Adjusted effect favoured TTS (*F*(1,17) = 42.23, *p* < 0.001); descriptive standard deviations internally inconsistent.",
     "⊕◯◯◯ Very low", "RoB (2); Ind (1); Imp (2)"],
    ["TTS-supported e-learning vs. baseline phases; reading/spelling and pronunciation [@s37]",
     "1 study; 1; SCED",
     "Higher performance described in the TTS phase; outcome scale undefined; effect not estimable.",
     "⊕◯◯◯ Very low", "RoB (2); Ind (1); Imp (2)"],
    ["Teacher instruction plus adaptive AI Braille tutor vs. teacher instruction alone; sessions to mastery of written Braille contractions [@s38]",
     "1 study; 10; SCED",
     "Mean 7.00 vs. 9.58 sessions to mastery (MD −2.58; CI not reported); different contraction sets per condition; described by the authors as a trend.",
     "⊕◯◯◯ Very low", "RoB (2); Ind (1); Imp (2)"],
]

TABLE3 = [
    ["Access and learning effects of TTS and STT are rarely separated; exposure is short",
     "Services cannot tell whether to expect immediate access benefits, durable skill gains, or both, which determines goals and review points.",
     "Parallel-group or cluster-randomised trials that prespecify an access or a learning question; factorial or dismantling designs separating technology from instruction; unaided outcomes at 6–12 months or later."],
    ["Effect modifiers identified only post hoc",
     "Reading profile, grade, attention, and comparator appear to determine who benefits.",
     "Prospectively specified, adequately powered moderator analyses; pooling of individual participant data from within-participant experiments."],
    ["Weak single-case designs (7 of 13 rated reports did not meet WWC standards)",
     "Single-case research is essential for low-incidence populations but is credible only when design standards are met.",
     "At least three demonstrations of effect; adequate, stable baselines; randomised phase starts where feasible; independent outcome measurement; open phase-level data."],
    ["No learner-level evaluation of automatic classroom captioning; no deaf or hard-of-hearing participants",
     "Captioning is widely deployed on the basis of word-error rates measured outside classrooms.",
     "Within-participant or cluster trials in real classrooms measuring access to teacher and peer talk, comprehension, listening effort, and participation across acoustic conditions and languages."],
    ["One small study of an adaptive Braille tutor; no learner-level OBR study",
     "Braille literacy is foundational for blind learners, and shortages of specialist teachers drive interest in technology.",
     "Comparative studies co-designed with teachers of students with visual impairments, measuring Braille reading fluency, comprehension, and independence rather than isolated contractions."],
    ["No data on participation, burden, stigma, privacy, harms, abandonment, or cost; 24 of 27 reports from high-income countries",
     "These outcomes determine continued use and value for money in health, rehabilitation, and education services.",
     "A core outcome set mapped to ICF-CY; mixed-methods and economic evaluations; studies in low- and middle-income settings."],
]

ABBREV = [("ADHD", "Attention-deficit/hyperactivity disorder"), ("AI", "Artificial intelligence"),
          ("CI", "Confidence interval"),
          ("GRADE", "Grading of Recommendations Assessment, Development and Evaluation"),
          ("ICF-CY", "International Classification of Functioning, Disability and Health for Children and Youth"),
          ("ID", "Intellectual disability"), ("LD", "Learning disability"), ("MD", "Mean difference"),
          ("NRS", "Non-randomised or counterbalanced within-participant study"),
          ("OBR", "Optical Braille recognition"), ("PRESS", "Peer Review of Electronic Search Strategies"),
          ("PRISMA", "Preferred Reporting Items for Systematic Reviews and Meta-Analyses"),
          ("RCT", "Randomised controlled trial"),
          ("RoB 2", "Cochrane risk-of-bias tool for randomised trials, version 2"),
          ("ROBINS-I", "Risk Of Bias In Non-randomised Studies of Interventions"),
          ("SCED", "Single-case experimental design"), ("STT", "Speech-to-text"),
          ("SWiM", "Synthesis Without Meta-analysis"), ("TTS", "Text-to-speech"), ("WWC", "What Works Clearinghouse")]


# ------------------------------------------------------------------ template handling
def prepare_template(doc):
    body = doc.element.body
    children = list(body)
    # keep: article type, title, authors, front-matter frame table, affiliations (5 paragraphs)
    keep_n = 9
    kept = children[:keep_n]
    sect_elems = [c for c in body.iter(qn("w:sectPr"))]
    sect_first = copy.deepcopy(sect_elems[0])       # portrait, first-page header, footer
    sect_land = copy.deepcopy(sect_elems[1])        # landscape
    sect_port = copy.deepcopy(sect_elems[2])        # portrait
    final = children[-1]
    assert final.tag == qn("w:sectPr")
    for c in children[keep_n:-1]:
        body.remove(c)
    # final section: continue page numbering, no distinct first page
    for tag in ("w:pgNumType", "w:titlePg", "w:bidi"):
        for el in final.findall(qn(tag)):
            final.remove(el)
    # sanity checks on kept elements
    texts = ["".join(t.text or "" for t in k.iter(qn("w:t"))) for k in kept]
    assert texts[0].strip() == "Systematic Review", texts[0]
    assert "Speech and Braille" in texts[1]
    assert "Correspondence" in texts[8], texts[8]
    # MDPI affiliation format: no position titles, no trailing full stops
    for k in kept[4:8]:
        ts = list(k.iter(qn("w:t")))
        for t in ts:
            if t.text:
                t.text = t.text.replace("Dean, Makkah National College", "Makkah National College")
                t.text = t.text.replace("Riyadh, Saudi Arabia.", "Riyadh, Saudi Arabia")
    # remove the Children journal logo from the first-page header (keep the MDPI logo)
    removed = 0
    for ref in sect_first.findall(qn("w:headerReference")):
        part = doc.part.related_parts[ref.get(qn("r:id"))]
        hdr = part.element
        for drawing in list(hdr.iter(qn("w:drawing"))):
            names = [d.get("descr", "") for d in drawing.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")]
            if any("Children journal logo" in n for n in names):
                run = drawing.getparent()
                run.getparent().remove(run)
                removed += 1
    assert removed == 1, "journal logo not found"
    return kept[1], sect_first, sect_land, sect_port


def set_title(par_el, doc, text):
    from docx.text.paragraph import Paragraph
    p = Paragraph(par_el, doc._body)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.add_run(text)


# ------------------------------------------------------------------ build
def build():
    numbers = assign()
    with open(os.path.join(HERE, "data", "ref_numbers.json"), "w") as fh:
        json.dump(numbers, fh, indent=1)
    inline = Inline(numbers, INCLUDED_ORDER)
    studies = load_csv("included_studies.csv")
    ed = load_csv("effect_direction.csv")

    doc = Document(TEMPLATE)
    title_el, sect_first, sect_land, sect_port = prepare_template(doc)
    first_portrait_used = False

    def para(kind, text, bold=False):
        p = doc.add_paragraph(style=STYLE[kind])
        add_runs(p, inline.parse(text, {"b": True} if bold else None))
        if kind == "tabcap":  # keep a table caption on the same page as its table
            p.paragraph_format.keep_with_next = True
        return p

    src = [l for l in open(SRC, encoding="utf-8").read().splitlines() if not l.startswith("%")]
    blocks, buf = [], []
    for line in src + [""]:
        if not line.strip():
            if buf:
                blocks.append(" ".join(buf))
                buf = []
            continue
        if line.startswith("::"):
            if buf:
                blocks.append(" ".join(buf))
                buf = []
            blocks.append(line)
        else:
            buf.append(line.strip())

    abstract_words = 0
    for b in blocks:
        if not b.startswith("::"):
            para("p", b)
            continue
        m = re.match(r"::(\w+)\s*(.*)", b)
        kind, arg = m.group(1), m.group(2)
        if kind == "title":
            set_title(title_el, doc, arg)
        elif kind in ("h1", "h2", "h3", "bullet", "numbered", "keywords", "figcap", "tabcap", "tabfoot",
                      "back", "pni", "p"):
            para(kind, arg)
        elif kind == "abstract_head":
            para("abstract_head", arg, bold=True)
        elif kind == "hq":
            para("hq", arg, bold=True)
        elif kind == "abstract":
            para("abstract", arg)
            abstract_words = len(inline.plain(arg).split())
        elif kind == "line":
            doc.add_paragraph(style="MDPI_1.9_line")
        elif kind == "figure":
            path, width = arg.split("|")
            p = doc.add_paragraph(style="MDPI_5.2_figure")
            p.add_run().add_picture(os.path.join(OUTDIR, path), width=Inches(float(width)))
        elif kind == "end_section":
            if arg == "portrait":
                section_break_paragraph(doc, sect_port if first_portrait_used else sect_first)
                first_portrait_used = True
            else:
                section_break_paragraph(doc, sect_land)
        elif kind == "table":
            if arg == "table1":
                rows, grp = table1_rows(studies, ed)
                three_line_table(doc, ["Study [ref.]; country", "Design", "Learners with disability (n); population",
                                       "Technology vs. comparator", "Outcome domains (A/L)", "Appraisal"],
                                 rows, [2300, 1900, 2600, 3300, 3300, 1900], inline, group_rows=grp)
            elif arg == "table2":
                three_line_table(doc, ["Evidence body: technology vs. comparator; outcome [ref.]",
                                       "Studies; learners (n); designs", "What the evidence shows",
                                       "Certainty", "Reasons for downgrading (levels)"],
                                 TABLE2, [3700, 2100, 5700, 1300, 2500], inline, size=7.5)
            elif arg == "table3":
                three_line_table(doc, ["Evidence gap", "Why it matters for services", "Recommended designs and outcomes"],
                                 TABLE3, [2700, 3300, 4400], inline)
            elif arg == "abbrev":
                t = three_line_table(doc, ["Abbreviation", "Definition"], [[a, d] for a, d in ABBREV],
                                     [1500, 6300], inline, size=9, indent_twips=2608)
            else:
                raise ValueError(arg)
        elif kind == "references":
            for key, n in sorted(numbers.items(), key=lambda kv: kv[1]):
                p = doc.add_paragraph(style="MDPI_8.1_references")
                add_runs(p, [(REFS[key], {})])
        else:
            raise ValueError(f"Unknown directive {kind}")

    cp = doc.core_properties
    cp.title = "Assistive Speech and Braille Technologies for Children and Adolescents with Disabilities: A Systematic Review"
    cp.subject = "Systematic review submitted to Healthcare (MDPI)"
    cp.keywords = "assistive technology; text-to-speech; speech-to-text; Braille; systematic review"
    cp.comments = ""
    normalise(doc)
    doc.save(OUT)
    print("saved", OUT)
    print("abstract words:", abstract_words)
    return abstract_words


if __name__ == "__main__":
    build()
