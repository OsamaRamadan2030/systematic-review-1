import json, re
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from studies import STUDIES

T = json.load(open('/home/user/build/supp_tables.json'))
TITLE = re.search(r'^@TITLE (.*)$', open('/home/user/build/manuscript.txt').read(), re.M).group(1)
FONT = 'Palatino Linotype'

def clean(s):
    R = [
        ('full text unavailable during revision', 'full text not accessible for verification'),
        ('full methods unavailable during revision', 'full methods not accessible for verification'),
        ('full study unavailable during revision', 'full study not accessible for verification'),
        ('final full text unavailable during revision', 'final full text not accessible for verification'),
        ('The original author extraction reports', 'The extracted data record'),
        ('The original author extraction gives', 'The extracted data give'),
        ('These detailed statistics were not independently verified against full text during revision.', 'These statistics could not be verified against the full text.'),
        ('but these exact phase means were not independently verified against full text during revision.', 'but these phase means could not be verified against the full text.'),
        ('The supplied extraction flags', 'The extraction record flags'),
        ('A published corrigendum is listed with the source references; its existence is not treated as a new study.', 'A published corrigendum altered signs in one table (Corrigendum, 2021).'),
        (' Final A 2 has four observations; an inadequate-final-baseline assertion is not retained.', ' The final baseline phase has four observations.'),
        ('Cohen\'s d', 'Cohen’s d'),
    ]
    for a, b in R:
        s = s.replace(a, b)
    # drop sentences that refer to formal ratings not reported in this review
    s = re.sub(r'(?<=[.;]) [^.;]*formal (?:design-standard|high-risk)[^.]*\.', '', s)
    s = re.sub(r'(?<=[.;]) [^.;]*no formal design-standard rating[^.]*\.', '', s)
    # leading zeros for decimals (MDPI style)
    s = re.sub(r'(?<![\w.])\.(\d)', r'0.\1', s)
    s = s.replace('A 1, B 1, A 2, B 2', 'A1, B1, A2, B2').replace('25 th', '25th')
    return s.strip()

d = docx.Document()
st = d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
sec = d.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ('left_margin', 'right_margin'): setattr(sec, m, Cm(1.8))
sec.top_margin = sec.bottom_margin = Cm(1.6)
# page numbers
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for t, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    r = fp.add_run()
    if t:
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t); r._r.append(e)
    else:
        e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._r.append(e)

def add_runs(p, text, size=None, bold=False):
    for tok in re.split(r'(\*[^*]+\*)', text):
        if not tok: continue
        it = tok.startswith('*') and tok.endswith('*')
        r = p.add_run(tok[1:-1] if it else tok); r.italic = it; r.bold = bold
        if size: r.font.size = Pt(size)

def P(text, size=10, bold=False, align=None, after=6, before=0):
    p = d.add_paragraph(); add_runs(p, text, size, bold)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if align: p.alignment = align
    return p

def H(text):
    p = P(text, 11, True, before=10, after=6); p.paragraph_format.keep_with_next = True

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); tcPr.append(s)

def table(header, rows, widths, size=8):
    t = d.add_table(rows=1, cols=len(header)); t.style = 'Table Grid'; t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ''; add_runs(c.paragraphs[0], h, size, True); shade(c, 'DCE6F0')
    trPr = t.rows[0]._tr.get_or_add_trPr(); e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)
    for row in rows:
        if isinstance(row, str):
            cells = t.add_row().cells; m = cells[0].merge(cells[-1]); m.text = ''
            add_runs(m.paragraphs[0], row, size, True); shade(m, 'EDEDED'); continue
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            for k, line in enumerate(v.split('\n')):
                p = cells[i].paragraphs[0] if k == 0 else cells[i].add_paragraph()
                add_runs(p, line, size, bold=(i == 0 and k == 0))
                p.paragraph_format.space_after = Pt(1)
    for r in t.rows:
        trPr = r._tr.get_or_add_trPr(); e = OxmlElement('w:cantSplit'); trPr.append(e)
        for i, c in enumerate(r.cells):
            c.width = Cm(widths[i])
    tbl = t._tbl
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[i] * 567)))
    tblPr = tbl.tblPr
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(int(sum(widths) * 567))); w.set(qn('w:type'), 'dxa')
    for old in tblPr.findall(qn('w:tblW')): tblPr.remove(old)
    tblPr.append(w)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    return t

P('Supplementary Materials', 14, True, after=4)
P(TITLE, 12, True, after=4)
P('Ashit Kumar Dutta and Nasser Ali Aljarallah', 10, after=10)
P('Contents: Table S1, search strategies, dates, and yields; Table S2, characteristics and findings of included reports; Table S3, structured critical appraisal of included reports; Table S4, changes to the registered methods; Table S5, linked reports and potentially overlapping cohorts; Table S6, reports excluded after assessment or awaiting classification. Citations refer to the reference list of the main article. Numerical values are reproduced as reported in each source; missing values were not imputed. Findings verified only against an abstract, preview, figure, or linked primary document are provisional and are identified in Tables S2 and S3.', 9.5, after=8)

# ---------- S1 ----------
H('Table S1. Search strategies, dates, and yields.')
s1 = [[clean(r[0]).replace('\nRecords retrieved: n = ', '\nRecords: '), r[1]] for r in T[0][1:]]
s1[-1][1] = s1[-1][1]
s1.append(['ERIC (Phase 2)\n2 October 2026\nRecords: 267',
           '("text-to-speech" OR "speech-to-text" OR "speech recognition" OR "automatic captioning" OR "Braille tutor" OR "optical Braille") AND (disabilit* OR dyslex* OR blind* OR deaf* OR "reading difficulties" OR "writing difficulties" OR ADHD)\nExecuted through the ERIC application programming interface (https://api.ies.ed.gov/eric/; format = json; rows = 1000; start = 0). All 267 returned records were screened. No date, language, or publication-type filter was applied; eligibility restrictions, including the 23 July 2026 publication cutoff, were applied at screening.'])
table(['Source, date, and yield', 'Strategy as executed'], s1, [4.2, 21.9], 8)
P('Phase 1 other methods (July 2026): Google Scholar, first 200 results (strategy above); backward and forward citation searching, 43 records; publisher website searching, 19 records; verification searching, 8 records. The seed articles, websites, and dates for these supplementary methods were not fully recorded. The PsycInfo host platform and the Web of Science collections searched were not recorded, and the database exports were not retained. Phase 2 accounting: 2 duplicates within ERIC; 15 records matching reports included in Phase 1; 219 records excluded at title and abstract screening; 31 reports assessed; together with 6 candidates from targeted and citation searching, 37 unique candidate reports were assessed (21 included, 7 excluded, 9 awaiting classification; Table S6). Targeted searching of publishers, repositories, and related reports is recorded per report in Tables S2, S3, and S6.', 8.5, before=4)

# ---------- S2 ----------
H('Table S2. Characteristics and findings of included reports (n = 47).')
P('Reports are listed alphabetically. Enrolled numbers are not necessarily analyzed denominators, and no total participant count is calculated because several cohorts overlap (Table S5). Phase: 1, July 2026 searches; 2, October 2026 supplementary searches. Verification: full text, or limited (abstract, preview, figure, or linked primary document).', 8.5)
meta = {s[0]: s for s in STUDIES}
def key_for(label):
    lab = label.split('\n')[0]
    lab = re.sub(r';.*', '', lab).strip()
    m = re.match(r'(.+?) \((\d{4})\)', lab)
    name, yr = m.group(1), m.group(2)
    name = name.replace(' and ', ' & ')
    for k in meta:
        if k.startswith(name) and k.endswith(yr): return k
    raise KeyError(label)
s2 = []
for r in T[1][1:]:
    k = key_for(r[0]); s = meta[k]
    first, _, design = r[0].partition('\n')
    head = f'{first}\nPhase {s[9]}; {"full text" if s[10] == "F" else "limited"}'
    s2.append([clean(head), clean(design), clean(r[1]), clean(r[2]).replace(' versus ', '\nversus '), clean(r[3])])
s2.sort(key=lambda x: x[0].replace('ä', 'a').replace('ü', 'u'))
table(['Report', 'Design', 'Participants', 'Technology and comparator', 'Outcomes and findings'], s2, [3.4, 3.6, 4.6, 5.4, 9.1], 7.5)

# ---------- S3 ----------
H('Table S3. Structured critical appraisal of included reports.')
P('Domains considered: allocation or phase control; confounding and co-interventions; missing observations; outcome measurement; reporting consistency; applicability; and depth of source access. This is a structured narrative appraisal; formal RoB 2, ROBINS-I, or What Works Clearinghouse ratings are not reported (main text, Section 2.7), and no cross-tool score or GRADE rating is assigned. Precision and applicability are distinguished from risk of bias.', 8.5)
s3 = [[clean(r[0]), clean(r[1]), clean(r[2])] for r in T[2][1:]]
s3.sort(key=lambda x: x[0].replace('ä', 'a').replace('ü', 'u'))
table(['Report', 'Source accessed for verification', 'Principal limitations'], s3, [3.8, 7.6, 14.7], 7.5)

# ---------- S4 ----------
H('Table S4. Changes to the registered methods.')
s4 = [
 ['Age eligibility', 'Explicit primary or secondary school enrolment accepted when ages were incompletely reported; stated age limits otherwise applied, and missing ages were not imputed.', 'Several eligible reports gave grade levels but not ages.', 'Phase 2'],
 ['Comparative designs', 'Within-participant comparisons eligible without counterbalancing (order effects appraised). Observational use/nonuse analyses eligible when eligible-subgroup outcomes were separable; interpreted as associations.', 'To retain relevant comparative evidence while preventing causal over-interpretation.', 'Phase 2'],
 ['Technology contrast', 'Packages eligible only when an eligible technology was central to the contrast; comparisons that held the technology constant and varied only instruction excluded. Prerecorded human speech remained out of scope; this led to exclusion of Bonifacci et al. (2022).', 'Operational clarification of the registered intervention criterion.', 'Phase 2'],
 ['Outcome classification', 'Outcomes classified as assisted, unaided, proximal, or unclear according to the measurement condition.', 'To avoid conflating access (accommodation) effects with learning effects.', 'After Phase 1 extraction'],
 ['Searches', 'ERIC and targeted searches added on 2 October 2026, retaining the 23 July 2026 publication cutoff; shown as a separate pathway in Figure 1.', 'Education-specific coverage of the Phase 1 searches was limited.', 'Phase 2'],
 ['Report linkage', 'Companion reports and shared cohorts identified explicitly; no total participant count calculated.', 'To avoid double counting of participants and replications.', 'Phase 2'],
 ['Critical appraisal', 'Planned RoB 2, ROBINS-I, and What Works Clearinghouse ratings replaced by structured critical appraisal (Table S3); the sensitivity analysis based on those ratings was removed.', 'Documentation supporting several overall judgments was incomplete or did not support the rating; ratings from different tools are not comparable.', 'Phase 2'],
 ['Certainty and presentation', 'Formal GRADE ratings and effect-direction (vote-count) displays replaced by qualitative, outcome-specific confidence statements (Section 2.8; Table 2).', 'Outcomes with different measurement conditions had been combined; vote counting ignores precision and design.', 'Phase 2'],
]
table(['Item', 'Change', 'Reason', 'Timing'], s4, [3.4, 11.0, 8.5, 3.2], 8)

# ---------- S5 ----------
H('Table S5. Linked reports and potentially overlapping cohorts.')
s5 = [
 ['Raskind and Higgins (1999); Higgins and Raskind (2000)', 'The 2000 report includes the 39 participants of the 1999 report and adds 13 continuous-speech users.', 'Distinct comparisons retained; not treated as independent replication.'],
 ['Keelor et al. (2018, 2023)', 'Linked to the same 29-child experiment through the doctoral dissertation and matching conditions (provisional).', 'The 2018 report contributes predictor analyses only; no independent cohort.'],
 ['Keelor et al. (2020) and the other Keelor reports', 'Possible overlap unresolved.', 'Not counted as independent replication; no participant sums.'],
 ['Wei (2024); Ogut et al. (2025)', 'Both analyze 2017 NAEP grade 8 mathematics data.', 'Rounded samples and student–item observations are not additive.'],
 ['Fälth, Björklund, et al. (2025); Fälth, Nilsson, et al. (2025); Sand et al. (2025); Svensson et al. (2026)', 'Shared Swedish research program and ethics approval.', 'Independence of participants not established.'],
 ['Meyer and Bouck (2014, 2017)', 'Different sample sizes and designs; independence unresolved.', 'Interpreted as potentially related.'],
 ['Garrett et al. (2011); Keelor et al. (2018, 2023); Park et al. (2017); Noakes et al. (2019)', 'Related dissertations or precursor reports used to check eligibility, design, or linkage.', 'Their numerical results were not substituted for journal results and they are not counted as additional reports.'],
]
table(['Reports', 'Relationship', 'Handling in the synthesis'], s5, [8.0, 9.0, 9.1], 8)

# ---------- S6 ----------
H('Table S6. Reports excluded after assessment or awaiting classification.')
s6 = [
 'Excluded on re-checking of Phase 1 reports',
 ['Bonifacci et al. (2022)', 'Excluded (intervention)', 'Experimental audio consisted of words recorded in a human voice and presented to simulate synthesized speech (report Section 5.2.6); this does not constitute automatic TTS and falls under the prerecorded-speech exclusion.'],
 ['Schiavo et al. (2021)', 'Excluded (intervention)', 'Read-aloud used prerecorded human speech.'],
 ['Stodden et al. (2012)', 'Excluded (publication type)', 'Conference proceedings paper. The related journal report (Park et al., 2017) is included provisionally.'],
 'Phase 2 candidates: excluded (n = 7)',
 ['Argyropoulos et al. (2009)', 'Excluded (population)', 'All 60 participants were university students aged 20–25 years; no eligible child or adolescent subgroup.'],
 ['Daley et al. (2020)', 'Excluded (population)', 'Of 315 grade 6–8 remedial readers, 177 had IEPs; TTS–comprehension models adjusted for IEP status but did not report IEP-specific TTS estimates, so results could not be attributed to eligible learners.'],
 ['Fasting and Lyster (2005)', 'Excluded (population)', 'Teacher-selected struggling readers and spellers without diagnosis, special-education designation, or standardized cutoff; no qualifying subgroup separately extractable.'],
 ['Felix et al. (2017)', 'Excluded (intervention)', 'System used recorded audio with pictographs, pronunciation recognition, and touchscreen handwriting recognition; not dictation, captioning, synthesized TTS, or Braille tutoring or recognition.'],
 ['Haug and Klein (2018)', 'Excluded (population)', 'Unselected general grade 5 classes; no separately identifiable eligible subgroup.'],
 ['Kambouri et al. (2023)', 'Excluded (design)', 'Uncontrolled group pre–post evaluation; handwriting–STT contrasts compared baseline handwriting with post-intervention STT, without a contemporaneous modality comparison.'],
 ['Sulaimon et al. (2025)', 'Excluded (intervention contrast)', 'TTS listening and STT responding were present in all phases; the manipulated component was self-questioning instruction, not the technology.'],
 'Phase 2 candidates: awaiting classification (n = 9)',
 ['Douglas et al. (2009)', 'Awaiting classification', 'Abstract summarizes six single-case experiments with supported eText but does not separate synthesized TTS from recorded-voice conditions or give eligible denominators.'],
 ['Edeh et al. (2025)', 'Awaiting classification', 'NAEP 2019 analysis (disability defined as IEP/504 plan); the abstract does not establish an eligible TTS versus non-TTS comparison for a learner outcome.'],
 ['Higgins and Raskind (2004)', 'Awaiting classification', 'Population, comparator, and outcomes appear eligible, but the exact technology function and relationship to the authors’ earlier cohorts require the full methods.'],
 ['Leong (1995)', 'Awaiting classification', 'Abstract does not establish diagnostic status or a standardized cutoff; reading-ability groups are mixed.'],
 ['Quinlan (2004)', 'Awaiting classification', 'Design and age appear eligible, but the abstract does not establish a formally identified disability, IEP, or standardized threshold.'],
 ['Schmitt et al. (2011)', 'Awaiting classification', 'Remedial readers; formal disability, IEP, or a standardized selection cutoff not established in the abstract.'],
 ['Wei (2025)', 'Awaiting classification', 'NAEP 2017 grade 8 observational analysis involving students with disabilities; separate TTS–performance results and denominators not reported in the abstract.'],
 ['Wei (2026)', 'Awaiting classification', 'NAEP 2017 analysis of 28,090 students stratified by proficiency, not disability; an eligible subgroup and separable results are not established.'],
 ['Witmer and Bouck (2023)', 'Awaiting classification', 'Abstract reports tool-use predictors and performance associations but not a separately extractable disability sample or TTS-specific contrast.'],
 'Other',
 ['Alsalman (2026)', 'Not eligible (date)', 'Published online on 22 September 2026, after the 23 July 2026 cutoff; the accessible abstract does not establish an adaptive Braille function.'],
 ['Phase 1 full-text exclusions (n = 101)', 'Excluded', 'Reasons are summarized in Figure 1 of the main article; the citation-level list was not retained.'],
]
table(['Report', 'Decision', 'Reason'], s6, [5.0, 4.0, 17.1], 8)

d.save('/home/user/build/out/Supplementary_Materials.docx')
print('ok')
